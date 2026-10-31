"""Genera los entregables del art. 13 de la Directiva de originalidad.

La RR N.° 000947-2026-R, art. 13 apartado A, exige que la tesis se entregue en
**dos archivos .docx** con nombres exactos, dentro de una estructura de carpetas
concreta, y que las tablas y fórmulas estén como texto y no como imágenes. Todo
eso se genera aquí; el tesista no tiene que recordar ninguna de esas reglas.

Lo que produce:

    entregables/
      TEXTO/
        Apellidos Nombres - Desarrollo Título - Año.docx
        Apellidos Nombres - Protocolo y Complemento Título - Año.docx
      DOCUMENTO/
        checklist_expediente.md
        declaracion_jurada_pids.md          (Anexo 4)
        declaracion_de_originalidad.md      (RENATI art. 8.2; va dentro del texto)
        declaracion_uso_de_ia.md            (voluntaria, ver más abajo)

Sobre la declaración de uso de IA: **ninguna norma leída la exige** (ver
`normativa/requisitos.md`, P-01). Se genera desde `registro/registro_ia.jsonl`
como medida voluntaria de transparencia, apoyada en el principio de transparencia
del art. 5 del Código de Ética. Se le presenta al tesista como tal, no como un
requisito.

Uso:
    uv run python -m herramientas.entregar --comprobar
    uv run python -m herramientas.entregar
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import pathlib
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field

import docx
import yaml

from herramientas import caratula, docx_numeracion, pdf, registro_ia, rutas
from herramientas.informe import Informe
from herramientas.perfiles import ErrorDePerfil, cargar_de, perfil_activo

RAIZ = pathlib.Path(__file__).resolve().parent.parent
DATOS = RAIZ / "entregables" / "datos_de_entrega.yaml"

#: Caracteres que Windows no admite en un nombre de archivo.
PROHIBIDOS = r'<>:"/\|?*'

PLANTILLA_DATOS = """\
# Datos para generar los entregables. Los completa /iniciar o el tesista.
# Se usan tal cual en los nombres de archivo que exige el art. 13, así que
# escríbelos como deben aparecer.

titulo: ""                 # título exacto de la tesis
anio: null                 # año de sustentación, p. ej. 2026

autores:                   # hasta tres (RR 00744-R-20, §7.2 a)
  - apellidos: ""
    nombres: ""
    dni: ""
    objetivos: []          # obligatorio si la tesis es grupal (§7.2 b)

asesor:
  nombres_y_apellidos: ""
  correo: ""               # institucional
  orcid: ""                # OBLIGATORIO para el deposito (Anexo 4)
  dni: ""                  # OBLIGATORIO para el deposito (Anexo 4)

facultad: ""
escuela: ""
titulo_profesional: ""     # "Economista", "Licenciado en Educación", …
grupo_de_investigacion: ""
fuente_de_financiamiento: ""
disciplina_ocde: ""
titulo_en_ingles: ""
linea_de_investigacion: ""   # algunas facultades la piden en la carátula

# Metadata complementaria del depósito en Cybertesis (RR N.° 06893-R-18).
# Va en una hoja después de la carátula.
orcid_autor: ""            # opcional: la directiva dice que no todos lo tienen
ubicacion_geografica: ""   # localidades y/o coordenadas donde se hizo el estudio
anios_de_la_investigacion: ""   # año o rango, p. ej. "2024-2025"
domicilio_para_el_deposito: ""  # lo pide el Anexo 3.1 de la autorización
correo_institucional: ""   # tu correo @unmsm.edu.pe, lo pide el Anexo 3.1
telefono: ""               # lo pide el Anexo 3.1
tipo_de_documento: "Tesis" # Tesis, Trabajo de investigación…
"""


@dataclass
class Autor:
    apellidos: str
    nombres: str
    dni: str = ""
    objetivos: list[str] = field(default_factory=list)

    @property
    def nombre_para_archivo(self) -> str:
        return f"{self.apellidos} {self.nombres}".strip()


def _limpiar(texto: str) -> str:
    """Quita lo que Windows no admite en un nombre de archivo."""
    limpio = "".join(" " if c in PROHIBIDOS else c for c in texto)
    return re.sub(r"\s+", " ", limpio).strip()


def nombre_archivo(autores: list[Autor], rotulo: str, titulo: str, anio: int,
                   extension: str = "docx") -> str:
    """Nombre exacto del art. 13 A.3.

    La directiva ejemplifica:
        Julca Pereda Dante Jorge - Desarrollo Competencia digitales en personas
        adultas - 2025.docx
    Obsérvese que entre el rótulo y el título no lleva guion.
    """
    if not autores:
        raise ValueError("no hay autores para nombrar el archivo")
    quien = autores[0].nombre_para_archivo
    if len(autores) > 1:
        quien = " y ".join(a.nombre_para_archivo for a in autores)
    return _limpiar(f"{quien} - {rotulo} {titulo} - {anio}") + f".{extension}"


def cargar_datos(ruta: pathlib.Path = DATOS) -> dict:
    if not ruta.exists():
        raise FileNotFoundError(ruta)
    return yaml.safe_load(ruta.read_text(encoding="utf8")) or {}


def leer_autores(datos: dict) -> list[Autor]:
    salida: list[Autor] = []
    for a in datos.get("autores") or []:
        if not isinstance(a, dict):
            continue
        salida.append(Autor(
            apellidos=(a.get("apellidos") or "").strip(),
            nombres=(a.get("nombres") or "").strip(),
            dni=(a.get("dni") or "").strip(),
            objetivos=[str(o) for o in (a.get("objetivos") or [])],
        ))
    return salida


def comprobar_datos(datos: dict, inf: Informe) -> list[Autor]:
    autores = leer_autores(datos)
    if not autores:
        inf.error("T01", "no hay autores en `entregables/datos_de_entrega.yaml`")
    for i, a in enumerate(autores, 1):
        if not a.apellidos or not a.nombres:
            inf.error("T02", f"el autor {i} no tiene apellidos o nombres completos")
    if len(autores) > 3:
        inf.error("T03", f"hay {len(autores)} autores. La directiva general permite "
                         "un máximo de tres (RR 00744-R-20, §7.2 a)")
    if len(autores) > 1:
        sin_objetivos = [a.nombre_para_archivo for a in autores if not a.objetivos]
        if sin_objetivos:
            inf.error("T04", "en una tesis grupal cada participante debe tener "
                             "objetivos definidos (§7.2 b). Sin objetivos: "
                             + ", ".join(sin_objetivos))
    if not (datos.get("titulo") or "").strip():
        inf.error("T05", "falta el título de la tesis")
    anio = datos.get("anio")
    if not isinstance(anio, int):
        inf.error("T06", "falta el año de sustentación, o no es un número")
    return autores


# ---------------------------------------------------------------------------
# Comprobaciones sobre el .docx generado
# ---------------------------------------------------------------------------

#: Un EMU es 1/914400 de pulgada. Media pulgada de ancho es el corte entre un
#: icono y algo que de verdad podría ser una tabla fotografiada: el icono de un
#: callout de Quarto mide 152400 EMU, o sea 12 puntos.
MEDIA_PULGADA_EMU = 457200

RE_EXTENT = re.compile(r'<wp:extent\s+cx="(\d+)"\s+cy="(\d+)"')


def _una_sola_celda(tabla) -> bool:
    """¿Es la tabla de una celda con que pandoc envuelve cada figura?

    `len(tabla.columns)` revienta con `InvalidXmlError` si el `<w:tbl>` no
    trae `<w:tblGrid>`, cosa que pasa en las tablas que pandoc genera para los
    callouts. Se cuentan las celdas a mano.
    """
    if len(tabla.rows) != 1:
        return False
    return len(tabla.rows[0].cells) == 1


def _solo_iconos(tbl_xml: str) -> bool:
    """¿Todas las imágenes de esta tabla son del tamaño de un icono?"""
    medidas = RE_EXTENT.findall(tbl_xml)
    if not medidas:
        return False
    return all(int(cx) < MEDIA_PULGADA_EMU and int(cy) < MEDIA_PULGADA_EMU
               for cx, cy in medidas)


def _imagenes_sin_iconos(xml: str) -> int:
    """Cuántas imágenes del documento podrían ser una fórmula rasterizada.

    Ninguna fórmula cabe en 12 puntos, así que los iconos de los callouts no
    cuentan. Lo que no declara tamaño se cuenta igual: `<w:pict>` (VML) no
    trae `<wp:extent>`, y es preferible avisar de más a callar una imagen real.
    """
    pequenas = sum(1 for cx, cy in RE_EXTENT.findall(xml)
                   if int(cx) < MEDIA_PULGADA_EMU and int(cy) < MEDIA_PULGADA_EMU)
    total = xml.count("<w:drawing") + xml.count("<w:pict")
    return max(total - pequenas, 0)


def revisar_docx(ruta: pathlib.Path, inf: Informe) -> None:
    """Art. 13 A.6: tablas y fórmulas como texto, nunca como imágenes."""
    try:
        from docx import Document
    except ImportError:  # pragma: no cover
        inf.aviso("T10", "falta python-docx: no se pudo revisar el .docx")
        return

    doc = Document(str(ruta))
    xml = doc.element.body.xml

    imagenes_en_tabla = 0
    for tabla in doc.tables:
        tbl_xml = tabla._tbl.xml
        if "<w:drawing" not in tbl_xml and "<w:pict" not in tbl_xml:
            continue
        # Pandoc emite cada figura con su leyenda dentro de una tabla de una
        # sola celda. Eso es una figura, no una tabla de datos rasterizada, y
        # el art. 13 A.6 habla de lo segundo.
        if _una_sola_celda(tabla):
            continue
        # Los callouts de Quarto —`::: {.callout-note}`— salen como una tabla
        # de dos filas cuya primera celda lleva el icono del callout: una
        # imagen de 12 pt, no una tabla fotografiada.
        if _solo_iconos(tbl_xml):
            continue
        imagenes_en_tabla += 1
    if imagenes_en_tabla:
        inf.error("T11", f"{imagenes_en_tabla} tabla(s) contienen imágenes. El art. 13 "
                         "A.6 exige que las tablas estén como texto",
                  archivo=ruta)

    # Una ecuación correcta viaja como OMML (<m:oMath>). Si no hay ninguna y sí
    # hay imágenes sueltas, puede que las fórmulas se hayan rasterizado.
    tiene_omml = "<m:oMath" in xml
    imagenes = _imagenes_sin_iconos(xml)
    if imagenes and not tiene_omml:
        inf.aviso("T12", f"el documento tiene {imagenes} imágenes y ninguna ecuación "
                         "en formato de Word. Comprueba que ninguna fórmula esté "
                         "insertada como imagen (art. 13 A.6)",
                  archivo=ruta, imagenes=imagenes)

    if not doc.paragraphs:
        inf.error("T13", "el documento salió vacío", archivo=ruta)


# ---------------------------------------------------------------------------
# Documentos del expediente
# ---------------------------------------------------------------------------

def declaracion_uso_de_ia(raiz: pathlib.Path, autores: list[Autor], datos: dict) -> str:
    """Anexo voluntario, construido desde registro/registro_ia.jsonl."""
    acciones = registro_ia.leer(raiz)
    por_etapa = registro_ia.por_etapa(acciones)
    archivos = registro_ia.archivos_tocados(acciones)
    al_dia, motivo_registro = registro_ia.esta_al_dia(raiz)

    lineas = [
        "# Declaración de uso de inteligencia artificial generativa",
        "",
        "> **Esta declaración es voluntaria.** A la fecha de generarla, ninguna norma",
        "> de la UNMSM que se haya podido consultar exige declarar el uso de IA",
        "> generativa en una tesis de pregrado (ver `normativa/requisitos.md`, P-01).",
        "> Se emite por transparencia, apoyándose en el principio de transparencia del",
        "> art. 5 del Código de Ética en Investigación de la UNMSM",
        "> (RR N.° 012648-2023-R), y en que la Directiva de originalidad",
        "> (RR N.° 000947-2026-R, art. 5 q) admite la IA generativa «como apoyo,",
        "> siempre respetando la originalidad».",
        "",
        f"**Tesis:** {datos.get('titulo', '')}",
        "",
        "**Autor(es):** " + "; ".join(a.nombre_para_archivo for a in autores),
        "",
        "## Herramienta utilizada",
        "",
        "Claude Code (Anthropic), a través del repositorio `tesis-sanmarcos`.",
        "",
        "## Alcance de la participación",
        "",
        "La herramienta ejecutó trabajo operativo: exploración y preparación de datos,",
        "búsqueda y síntesis de literatura, análisis, redacción de borradores, formato",
        "y verificación. **Las decisiones de investigación las tomó el autor**, y están",
        "registradas una por una en `decisiones/`, con sus respuestas textuales.",
        "",
        "## Registro de acciones",
        "",
    ]
    if acciones:
        lineas += [
            f"- Acciones registradas: **{len(acciones)}**",
            f"- Archivos del repositorio que la herramienta escribió o editó: "
            f"**{len(archivos)}**",
            "",
            "| Etapa | Acciones |", "|---|---|",
        ]
        for etapa, n in por_etapa.items():
            lineas.append(f"| {etapa} | {n} |")
        if not al_dia:
            lineas += [
                "",
                f"> **El registro puede estar incompleto:** {motivo_registro}. Se "
                "declara aquí para no presentar como exhaustivo algo que no lo es.",
            ]
    else:
        lineas += [
            "No hay acciones registradas en `registro/registro_ia.jsonl`.",
            "",
            "> Eso no significa que la herramienta no se haya usado: significa que el",
            "> hook que lleva ese registro no estaba instalado. Se instala con",
            "> `uv run python -m herramientas.instalar_hooks`. Si vas a presentar esta",
            "> declaración, dilo tal cual en lugar de dejar entender que no hubo ayuda.",
        ]
    lineas += [
        "",
        "## Lo que no se hizo",
        "",
        "- No se simularon ni completaron datos, resultados ni testimonios.",
        "- No se citó ninguna fuente que no estuviera en la biblioteca Zotero del autor.",
        "- No se empleó ninguna técnica para reducir artificialmente la similitud ni",
        "  para evadir su detección.",
        "",
        f"Generada el {dt.date.today().isoformat()}.",
        "",
        "Firma del autor: ______________________",
    ]
    return "\n".join(lineas)


def declaracion_jurada_pids(autores: list[Autor], datos: dict) -> str:
    """Anexo 4 de la RR N.° 000947-2026-R, con los datos ya puestos."""
    lineas = [
        "# Anexo 4 — Declaración jurada para la evaluación de documentos con la PIDS",
        "",
        "> Formato tomado de la Directiva de originalidad y similitud de trabajos",
        "> académicos, de investigación y producción intelectual de la UNMSM,",
        "> aprobada por RR N.° 000947-2026-R, Anexo 4.",
        "",
    ]
    for a in autores:
        lineas += [
            f"Yo, **{a.nombres} {a.apellidos}**, identificado(a) con DNI "
            f"**{a.dni or '__________'}**, en mi calidad de autor(a) de la Tesis "
            f"titulada «{datos.get('titulo', '')}»,",
            "",
            "**DECLARO BAJO JURAMENTO** que autorizo la evaluación de mi documento a",
            "través de la Plataforma PIDS de la UNMSM, aceptando los resultados que se",
            "consignen en el informe correspondiente, así como las disposiciones,",
            "infracciones y sanciones que pudieran derivarse del análisis efectuado",
            "conforme a la normativa vigente.",
            "",
            "Me ratifico en lo expresado, firmando la presente declaración en señal de",
            "conformidad.",
            "",
            "Lima, ____ de ______________ del 20____",
            "",
            "Firma del autor(a): ______________________",
            f"DNI N.°: {a.dni or '__________'}",
            "",
            "---",
            "",
        ]
    return "\n".join(lineas)


def declaracion_de_originalidad(autores: list[Autor], datos: dict) -> str:
    """Página preliminar que exige el RENATI, art. 8.2, folio 11.

    «El texto completo del trabajo conducente a un grado o título deberá
    contener entre sus páginas preliminares el resultado del informe de
    similitud del trabajo o, en su defecto, una declaración de originalidad
    consignando los nombres completos del responsable de la asesoría y del
    autor.»

    Es la única exigencia del RENATI que cae dentro del documento y no en el
    trámite, y **no figura en los preliminares del Anexo 2A**. Lo normal en la
    UNMSM es cumplirla con el reporte de la PIDS; esta declaración es la vía
    alterna que la propia norma admite, para cuando la escuela no lo adjunta.
    """
    asesor = datos.get("asesor") or {}
    lineas = [
        "# Declaración de originalidad",
        "",
        "> Página preliminar exigida por el Reglamento del RENATI (texto integrado,",
        "> art. 8.2, folio 11), aprobado por RCD N.° 033-2016-SUNEDU/CD y modificado",
        "> por las RCD N.° 174-2019-SUNEDU/CD y N.° 084-2022-SUNEDU/CD. Va **dentro**",
        "> del texto completo que se deposita, no en la carpeta del expediente.",
        "",
        "> La norma admite dos formas de cumplirla: el resultado del informe de",
        "> similitud o, «en su defecto», esta declaración. Si tu escuela adjunta el",
        "> reporte de la PIDS a las páginas preliminares, con eso basta.",
        "",
    ]
    for a in autores:
        lineas += [
            f"Quien suscribe, **{a.nombres} {a.apellidos}**, identificado(a) con DNI "
            f"**{_dato(a.dni)}**, autor(a) del trabajo titulado "
            f"«{datos.get('titulo', '')}», declara que su contenido es original y que "
            "las fuentes empleadas están debidamente citadas y referenciadas.",
            "",
            "Lima, ____ de ______________ del 20____",
            "",
            f"Firma: ______________________   {a.nombres} {a.apellidos}",
            "",
        ]
    lineas += [
        "Responsable de la asesoría: "
        f"**{_dato(asesor.get('nombres_y_apellidos'))}**",
        "",
        "Firma: ______________________",
        "",
    ]
    return "\n".join(lineas)


def _dato(valor, etiqueta_vacia: str = "__________") -> str:
    texto = str(valor or "").strip()
    return texto or etiqueta_vacia


def autorizacion_cybertesis(autores: list[Autor], datos: dict) -> str:
    """Anexo 3.1: autorización para el depósito de la obra en CYBERTESIS.

    La directiva original es la RR N.° 06893-R-18, pero **su anexo fue
    modificado por la RR N.° 00584-R-19** del 05/02/2019. La versión vigente es
    la que reproduce el «Protocolo para el depósito en CYBERTESIS durante el
    estado de emergencia sanitaria COVID-19», que es el documento que la
    universidad publica con los formatos al día.

    La diferencia no es cosmética: la versión de 2019 pide un bloque de datos
    del autor y del documento que la de 2018 no tenía —correo institucional,
    teléfono, tipo de documento, grado o título al que se opta, asesores—.

    Lo que el tesista firma es una cesión de uso a título gratuito y a nivel
    mundial, así que conviene que lo lea antes de firmarlo y no después.
    """
    asesor = datos.get("asesor") or {}
    lineas = [
        "# Anexo 3.1 — Autorización para el depósito de obra en CYBERTESIS UNMSM",
        "",
        "> Formato vigente según la RR N.° 00584-R-19, que modificó el anexo de la",
        "> «Directiva para el depósito de documentos de Grados y Títulos en",
        "> CYBERTESIS» (RR N.° 06893-R-18).",
        "",
        "## Datos del autor (tesista)",
        "",
    ]
    for a in autores:
        lineas += [
            f"- **Nombre y apellidos:** {a.nombres} {a.apellidos}",
            f"- **DNI:** {_dato(a.dni)}  ·  Pasaporte o carné de extranjería "
            f"(solo extranjeros): __________",
            f"- **Email UNMSM:** {_dato(datos.get('correo_institucional'))}",
            f"- **Teléfono:** {_dato(datos.get('telefono'))}",
            "",
        ]

    lineas += [
        "## Datos del documento para grado o título",
        "",
        f"- **Tipo de documento:** {_dato(datos.get('tipo_de_documento'), 'Tesis')}",
        f"- **Nombre del grado o título:** "
        f"{_dato(datos.get('titulo_profesional'))}",
        f"- **Título del documento:** {_dato(datos.get('titulo'))}",
        f"- **Asesor(es):** "
        f"{_dato(asesor.get('nombres_y_apellidos') or _nombre_llano(asesor))}",
        f"- **Facultad, unidad académica:** {_dato(datos.get('facultad'))}"
        f"{' — ' + str(datos.get('escuela')) if datos.get('escuela') else ''}",
        "",
        "## Declaración",
        "",
        "Por medio del presente documento, afirmo y garantizo ser el legítimo, único",
        "y exclusivo titular de todos los derechos de propiedad intelectual sobre los",
        "documentos arriba mencionados, las obras, los contenidos, los productos y/o",
        "las creaciones en general (en adelante, los «Contenidos») que serán incluidos",
        "en el repositorio institucional CYBERTESIS de la Universidad Nacional Mayor",
        "de San Marcos (en adelante, la «Universidad»).",
        "",
        "También, doy seguridad de que los contenidos entregados se encuentran libres",
        "de toda contraseña, restricción o medida tecnológica de protección, con la",
        "finalidad de permitir que se puedan leer, descargar, reproducir, distribuir,",
        "imprimir, buscar y enlazar los textos completos, sin limitación alguna.",
        "",
        "Autorizo a la Universidad Nacional Mayor de San Marcos a publicar los",
        "Contenidos en los Repositorios Institucionales de la Universidad y, en",
        "consecuencia, en el Repositorio Nacional Digital de Ciencia, Tecnología e",
        "Innovación de Acceso Abierto, sobre la base de lo establecido en la Ley",
        "N.° 30035, sus normas reglamentarias, modificatorias, sustitutorias y",
        "conexas, y de acuerdo con las políticas de acceso abierto que la Universidad",
        "aplique en relación con sus Repositorios Institucionales. Autorizo",
        "expresamente toda consulta y uso de los Contenidos, por parte de cualquier",
        "persona, por el tiempo de duración de los derechos patrimoniales de autor y",
        "derechos conexos, a título gratuito y a nivel mundial.",
        "",
        "En consecuencia, la Universidad tendrá la posibilidad de divulgar y difundir",
        "los Contenidos, de manera total o parcial, sin limitación alguna y sin",
        "derecho a pago de contraprestación, remuneración ni regalía alguna a favor",
        "mío; en los medios, canales y plataformas que la Universidad y/o el Estado de",
        "la República del Perú determinen, a nivel mundial, sin restricción geográfica",
        "alguna y de manera indefinida, pudiendo crear y/o extraer los metadatos sobre",
        "los Contenidos, e incluir los Contenidos en los índices y buscadores que",
        "estimen necesarios para promover su difusión.",
        "",
        "Autorizo que los Contenidos sean puestos a disposición del público a través",
        "de la siguiente licencia:",
        "",
        "**Creative Commons Reconocimiento-NoComercial-CompartirIgual 4.0",
        "Internacional** — <https://creativecommons.org/licenses/by-nc-sa/4.0/>",
        "",
        "En señal de conformidad, suscribo el presente documento.",
        "",
    ]
    for a in autores:
        lineas += [
            "Firma: ______________________________________",
            f"Nombre: {a.nombres} {a.apellidos}",
            f"DNI: {_dato(a.dni)}  ·  Pasaporte o carné de extranjería: __________",
            f"Domicilio: {_dato(datos.get('domicilio_para_el_deposito'))}",
            "",
            "Fecha: ______ / ______ / ________",
            "",
            "---",
            "",
        ]
    lineas += [
        "## Si necesitas postergar la publicación",
        "",
        "Se puede, por **6, 12, 18 o 24 meses**, con una solicitud justificada. El",
        "formato está en `solicitud_postergacion.md`, que esta misma herramienta",
        "genera. Es el Anexo 3.2.",
    ]
    return "\n".join(lineas)


def _nombre_llano(asesor: dict) -> str:
    nombres = str(asesor.get("nombres", "")).strip()
    apellidos = str(asesor.get("apellidos", "")).strip()
    return f"{nombres} {apellidos}".strip()


def solicitud_postergacion(autores: list[Autor], datos: dict) -> str:
    """Anexo 3.2: postergar la publicación en CYBERTESIS.

    Es **opcional**, y existe por si el trabajo se va a patentar, se envió a una
    revista indexada o está bajo un convenio de confidencialidad. Se genera
    siempre porque es el tipo de cosa que nadie sabe que existe hasta que la
    necesita y ya es tarde.
    """
    a = autores[0] if autores else None
    lineas = [
        "# Anexo 3.2 — Solicitud de postergación de la publicación (OPCIONAL)",
        "",
        "> Formato del «Protocolo para el depósito en CYBERTESIS», anexo 3.2,",
        "> conforme a la RR N.° 06893-R-18 y su modificatoria RR N.° 00584-R-19.",
        "",
        "**Esto es opcional.** Solo tiene sentido si tu trabajo se va a patentar, lo",
        "enviaste a una revista indexada o está bajo un convenio de confidencialidad.",
        "Si no es tu caso, no lo presentes: la tesis se publica y ya está.",
        "",
        "---",
        "",
        "Señor Director",
        "Dirección General de Biblioteca y Publicaciones",
        "",
    ]
    nombre = f"{a.nombres} {a.apellidos}" if a else "__________"
    dni = _dato(a.dni) if a else "__________"
    lineas += [
        f"Yo, **{nombre}**, identificado(a) con DNI **{dni}**, de la "
        f"{_dato(datos.get('facultad'))}, en calidad de autor de la "
        f"{_dato(datos.get('tipo_de_documento'), 'tesis')} con título",
        "",
        f"«{_dato(datos.get('titulo'))}»",
        "",
        f"para obtener el título de {_dato(datos.get('titulo_profesional'))}, me "
        f"dirijo a usted para solicitar la postergación de la publicación del "
        f"documento mencionado en el repositorio institucional CYBERTESIS, por el "
        f"tiempo exacto de:",
        "",
        "- [ ] seis meses",
        "- [ ] doce meses",
        "- [ ] 18 meses",
        "- [ ] 24 meses",
        "",
        "justificado por:",
        "",
        "- [ ] trabajo propuesto para patente.",
        "- [ ] trabajo presentado para publicación en revista indexada.",
        "- [ ] trabajo en convenio de confidencialidad temporal.",
        "- [ ] Otro motivo: ______________________________________________",
        "",
        "Adjunto los documentos que sustentan el motivo antes señalado.",
        "",
        "Sin otro particular, saludo a usted atentamente.",
        "",
        "Firma: ______________________________________",
        f"Nombre y apellido: {nombre}",
        f"DNI: {dni}  ·  Pasaporte o carné de extranjería: __________",
        f"Correo de contacto: {_dato(datos.get('correo_institucional'))}",
    ]
    return "\n".join(lineas)


def metadata_complementaria(autores: list[Autor], datos: dict) -> str:
    """Anexo 4: la información complementaria del depósito.

    Versión vigente, la que fijó la RR N.° 00584-R-19 al modificar el anexo de
    la directiva de depósito. Tiene **tres campos marcados como obligatorios** y
    ninguno de los tres está en la información complementaria del Anexo 2A de la
    directiva general de tesis, que es de donde salía nuestra lista:

    - el ORCID del asesor,
    - el DNI del autor,
    - el DNI del asesor.

    El protocolo además pide este archivo **en Word**, no en PDF como el resto.
    """
    asesor = datos.get("asesor") or {}
    autor = autores[0] if autores else None

    filas = [
        ("Código ORCID del asesor(es)", asesor.get("orcid") or "", True),
        ("DNI del autor", (autor.dni if autor else "") or "", True),
        ("DNI del asesor", asesor.get("dni") or "", True),
        ("Código ORCID del autor", datos.get("orcid_autor") or "", False),
        ("Grupo de investigación", datos.get("grupo_de_investigacion") or "", False),
        ("Financiamiento", datos.get("fuente_de_financiamiento") or "", False),
        ("Ubicación geográfica donde se desarrolló la investigación "
         "(localidades y/o coordenadas)", datos.get("ubicacion_geografica") or "",
         False),
        ("Año o rango de años que la investigación abarcó",
         datos.get("anios_de_la_investigacion") or "", False),
    ]

    lineas = [
        "# Anexo 4 — Información complementaria",
        "",
        "> Formato vigente según la RR N.° 00584-R-19, que modificó el anexo de la",
        "> directiva de depósito en CYBERTESIS (RR N.° 06893-R-18).",
        "",
        "**Este archivo se entrega en Word**, a diferencia del resto del expediente,",
        "que va en PDF. Lo dice el numeral 5 del protocolo de depósito.",
        "",
        "| Campo | Valor | |",
        "|---|---|---|",
    ]
    for campo, valor, obligatorio in filas:
        marca = "**obligatorio**" if obligatorio else ""
        lineas.append(f"| {campo} | {valor or '**FALTA**'} | {marca} |")

    faltan_obligatorios = [c for c, v, o in filas if o and not v]
    faltan_otros = [c for c, v, o in filas if not o and not v]

    if faltan_obligatorios:
        lineas += [
            "",
            "## Falta lo obligatorio",
            "",
            "Sin estos tres campos, Biblioteca no deposita el documento:",
            "",
        ]
        lineas += [f"- {c}" for c in faltan_obligatorios]
        lineas += [
            "",
            "Los tres se completan en `entregables/datos_de_entrega.yaml`. El ORCID",
            "se saca de <https://orcid.org>; si tu asesor no lo tiene, lo crea en",
            "cinco minutos. **El DNI del asesor es el que más veces falta**, porque",
            "no aparece en ninguna otra parte del expediente.",
        ]
    if faltan_otros:
        lineas += ["", "## Falta completar", ""]
        lineas += [f"- {c}" for c in faltan_otros]
    return "\n".join(lineas)


def _metadata_en_word(destino: pathlib.Path, autores: list[Autor], datos: dict,
                      inf: Informe) -> None:
    """La información complementaria, además, en .docx.

    El protocolo de depósito es explícito: los archivos del expediente van en
    PDF «y el archivo de información complementaria en Word» (numeral 5). Es el
    único que se entrega en ese formato, así que es también el que más veces se
    manda mal.
    """
    try:
        doc = docx.Document()
        doc.add_heading("Información complementaria", level=1)
        doc.add_paragraph(
            "Anexo 4 de la directiva de depósito en CYBERTESIS, en su versión "
            "vigente (RR N.° 00584-R-19).")

        asesor = datos.get("asesor") or {}
        autor = autores[0] if autores else None
        filas = [
            ("Código ORCID del asesor(es)", asesor.get("orcid"), True),
            ("DNI del autor", autor.dni if autor else None, True),
            ("DNI del asesor", asesor.get("dni"), True),
            ("Código ORCID del autor", datos.get("orcid_autor"), False),
            ("Grupo de investigación", datos.get("grupo_de_investigacion"), False),
            ("Financiamiento", datos.get("fuente_de_financiamiento"), False),
            ("Ubicación geográfica donde se desarrolló la investigación "
             "(localidades y/o coordenadas geográficas)",
             datos.get("ubicacion_geografica"), False),
            ("Año o rango de años que la investigación abarcó",
             datos.get("anios_de_la_investigacion"), False),
        ]
        tabla = doc.add_table(rows=0, cols=2)
        tabla.style = "Table Grid"
        for campo, valor, obligatorio in filas:
            fila = tabla.add_row().cells
            fila[0].text = campo + (" [Es un dato obligatorio]" if obligatorio else "")
            fila[1].text = str(valor or "")

        ruta = destino / "informacion_complementaria.docx"
        doc.save(str(ruta))
        inf.info("T29", "información complementaria también en .docx, que es el "
                        "formato en que la pide el protocolo de depósito")
    except Exception as e:  # noqa: BLE001
        inf.aviso("T30", f"no se pudo escribir la información complementaria en "
                         f".docx: {e}. La versión en Markdown sí está")


#: Archivos de `datos/crudos/` que no son datos, sino papeleo de la carpeta.
#: `_procedencia.md` lo escribe **todo** proyecto que siga la etapa 1, también
#: uno que solo bajó series del Banco Mundial; si cuenta como transcripción, el
#: bloque de datos personales le sale a todo el mundo y deja de significar algo.
PAPELEO_DE_CRUDOS = ("_", "leeme", "readme")


def _hay_transcripciones(raiz) -> bool:
    """¿La tesis recogió datos de personas?

    Se mira si hay transcripciones en `datos/crudos/`. A quien trabajó solo con
    datos abiertos no le sirve de nada el bloque de datos personales, y un
    checklist con secciones que no le tocan se deja de leer entero.
    """
    if raiz is None:
        return False
    base = pathlib.Path(raiz) / "datos" / "crudos"
    if not base.is_dir():
        return False
    return any(
        r.suffix.lower() in {".txt", ".md", ".vtt", ".srt", ".docx"}
        and not r.name.lower().startswith(PAPELEO_DE_CRUDOS)
        for r in base.rglob("*") if r.is_file())


def checklist_expediente(perfil, autores: list[Autor], datos: dict) -> str:
    lineas = [
        "# Checklist del expediente",
        "",
        f"Perfil: `{perfil.id}` — {perfil.get('facultad')}",
        "",
        "## Lo que exige la Directiva de originalidad (art. 17)",
        "",
        "- [ ] **Certificado de Similitud** firmado por el asesor (Anexo 1).",
        "      Vence a los **90 días** de emitido.",
        "- [ ] Último **reporte de similitud** de la PIDS, dentro del rango permitido "
        "(0 % a 20 %).",
        "- [ ] **Historial de análisis** de la PIDS, si hubo varias revisiones.",
        "- [ ] **Anexo 4** firmado por cada autor (se generó en esta carpeta).",
        "- [ ] Comprobantes de pago de S/ 20 por cada revisión **a partir de la "
        "segunda** (art. 27). La primera es gratuita.",
        "",
        "## Archivos de entrega (art. 13)",
        "",
        "- [ ] Carpeta `TEXTO/` con los dos .docx.",
        "- [ ] Carpeta `DOCUMENTO/` nombrada según el tipo de documento.",
        "- [ ] Tablas y fórmulas como texto, no como imágenes.",
        "",
        "## Lo que exige el RENATI dentro del documento (art. 8.2)",
        "",
        "Es nacional, de la Sunedu, y es la única exigencia del RENATI que no es "
        "un trámite sino una página del texto. No está en los preliminares del "
        "Anexo 2A, así que es fácil que se pase:",
        "",
        "- [ ] El texto completo que se deposita lleva, **entre sus páginas "
        "preliminares**, el reporte de similitud de la PIDS o —«en su defecto», "
        "dice la norma— la **declaración de originalidad** con los nombres "
        "completos del asesor y del autor (se generó en esta carpeta).",
        "",
    ]

    if len(autores) > 1:
        lineas += ["## Tesis grupal (RR 00744-R-20, §7.2)", "",
                   f"- [ ] Son {len(autores)} autores (el máximo es tres).",
                   "- [ ] El problema es complejo o multidisciplinario.",
                   "- [ ] Cada participante tiene objetivos definidos:", ""]
        for a in autores:
            lineas.append(f"  - **{a.nombre_para_archivo}**")
            for o in a.objetivos:
                lineas.append(f"    - {o}")
        lineas.append("")

    etica = perfil.get("etica") or {}
    if etica.get("comite_requerido") is True:
        lineas += ["## Ética", "",
                   "- [ ] Aprobación del **Comité de Investigación** de la escuela.",
                   "- [ ] Aprobación de un **Comité de Ética en Investigación**.",
                   "- [ ] Resolución decanal de aprobación del proyecto.", ""]

    # Solo aparece si la tesis recogió datos de personas. A quien trabajó con
    # datos abiertos no le sirve de nada y sería ruido.
    if _hay_transcripciones(datos.get("_raiz")):
        lineas += [
            "## Datos personales (Ley N.° 29733 y DS N.° 016-2024-JUS)",
            "",
            "Tu tesis recogió datos de personas. Esto es información, no un "
            "trámite que te frene, pero conviene tenerlo resuelto:",
            "",
            "- [ ] **Consentimiento informado firmado** de cada participante. Si "
            "tocaste salud, vida afectiva o familiar, hábitos íntimos o "
            "afiliación sindical, son datos sensibles y el art. 8 del reglamento "
            "exige que sea por escrito. El art. 9 pone en ti la carga de probarlo.",
            "- [ ] **La tabla que relaciona códigos con nombres, destruida** si ya "
            "no la necesitas. Mientras exista, lo tuyo es disociación y no "
            "anonimización (art. 41.1), y el dato sigue siendo personal.",
            "- [ ] **Revisadas las citas del texto** para que nadie sea "
            "identificable por combinación de datos: oficio más barrio más edad "
            "más composición familiar puede señalar a una persona sin nombrarla.",
            "- [ ] Nada de `datos/crudos/` subido al repositorio remoto.",
            "",
            "Detalle con artículos en `normativa/requisitos.md`, R-28.",
            "",
        ]

    vigencia = perfil.get("vigencia_proyecto_meses")
    if isinstance(vigencia, int):
        lineas += ["## Vigencia del proyecto", "",
                   f"- [ ] El proyecto sigue vigente: caduca a los {vigencia} meses de "
                   "aprobado.", ""]

    if not perfil.verificado:
        lineas += ["## Atención", "",
                   f"- [ ] El perfil `{perfil.id}` está en estado `{perfil.estado}`: "
                   "la estructura usada es la general de la universidad (Anexos 2A "
                   "y 2B de la RR N.° 00744-R-20). Si tu escuela pide otra, "
                   "ajústala antes de imprimir.", ""]

    lineas += ["## Voluntario", "",
               "- [ ] Declaración de uso de IA (no exigida por norma; ver el archivo "
               "en esta carpeta).", ""]
    return "\n".join(lineas)


# ---------------------------------------------------------------------------
# Render
# ---------------------------------------------------------------------------

def _quarto() -> str | None:
    return shutil.which("quarto")


#: Prefijos de referencia cruzada de Quarto. `@fig-dispersion` no es una cita:
#: es un «ver figura 1». Si se cuela en el `nocite` del archivo de Complemento,
#: citeproc busca esa clave en el .bib, no la encuentra y la imprime en crudo.
PREFIJOS_CRUZADOS = ("fig-", "tbl-", "eq-", "sec-", "lst-", "thm-", "ref-")

RE_CITA = re.compile(r"(?<![\w@])@([A-Za-z][\w:.#$%&+?<>~/-]*)")


def claves_citadas(entradas: list[pathlib.Path]) -> list[str]:
    """Las claves de Zotero citadas en estos capítulos, en orden de aparición.

    Hacen falta porque la bibliografía se imprime en el archivo de «Protocolo y
    Complemento» (art. 13 A.2 b) mientras que las citas viven en el de
    Desarrollo. Sin esta lista, citeproc no tendría de dónde saber qué obras
    listar en un documento que no cita a ninguna.
    """
    vistas: dict[str, None] = {}
    for ruta in entradas:
        texto = ruta.read_text(encoding="utf8", errors="replace")
        for clave in RE_CITA.findall(texto):
            if not clave.startswith(PREFIJOS_CRUZADOS):
                vistas.setdefault(clave.rstrip(".,;:"), None)
    return list(vistas)


def _cabecera(raiz: pathlib.Path, meta: dict[str, str] | None = None) -> str:
    """El encabezado YAML del .qmd combinado.

    El archivo combinado no hereda el `_quarto.yml` del proyecto, porque no
    está en su lista de `render:`. Lo que ese archivo declara se repite aquí:

    - `lang`, para que las leyendas salgan en español («Figura 1», «Tabla 2»).
    - `csl`, para que la bibliografía siga el estilo que exige la facultad.
    """
    lineas = ["---", "lang: es"]
    if _bib(raiz):
        lineas.append(f'bibliography: "{_bib(raiz)}"')
    if _csl(raiz):
        lineas.append(f'csl: "{_csl(raiz)}"')
    for clave, valor in (meta or {}).items():
        lineas.append(f"{clave}: {valor}")
    lineas.append("---")
    return "\n".join(lineas)


def _render(entradas: list[pathlib.Path], destino: pathlib.Path,
            referencia: pathlib.Path, raiz: pathlib.Path,
            meta: dict[str, str] | None = None,
            cola: str = "") -> tuple[bool, str]:
    quarto = _quarto()
    if quarto is None:
        return False, "no se encontró quarto en el PATH"
    # El .qmd combinado se arma EN LA RAÍZ del proyecto, que es desde donde el
    # repositorio ejecuta el código (`execute-dir: project`). Así no hace falta
    # `--execute-dir`, que en Windows mata el kernel de Jupyter cuando apunta a
    # un directorio distinto del del archivo.
    #
    # El precio es que las rutas relativas de los capítulos —sus figuras, que
    # apuntan a `../resultados/…` desde `capitulos/`— dejan de ser válidas. Se
    # reubican una por una. Si no se hiciera, pandoc no encontraría la imagen,
    # la sustituiría por su texto alternativo y NO fallaría: saldría una tesis
    # sin figuras y sin aviso.
    combinado = raiz / f"_entrega_{abs(hash(destino.name)) % 10**8}.qmd"
    partes = [_reubicar_rutas(r.read_text(encoding="utf8", errors="replace"),
                              r.parent, raiz)
              for r in entradas]
    texto = _cabecera(raiz, meta) + "\n\n" + "\n\n".join(partes)
    if cola:
        texto += "\n\n" + cola + "\n"
    combinado.write_text(texto, encoding="utf8")

    salida = raiz / destino.name
    cmd = [quarto, "render", combinado.name, "--to", "docx",
           "--reference-doc", str(referencia.resolve()),
           "--output", destino.name]
    try:
        r = _correr_quarto(cmd, raiz)
    except Exception as e:  # noqa: BLE001
        return False, f"falló el render: {e}"
    finally:
        _limpiar_residuos(combinado)

    if r.returncode != 0 or not salida.exists():
        return False, _motivo_util(r.stderr, r.stdout) or "quarto no escribió el .docx"

    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        _mover(salida, destino)
    except OSError as e:
        return False, (f"no se pudo mover el documento a su carpeta: {e}. "
                       f"{_consejo_ruta_larga(destino)}")

    perdidas = [ln.split("Could not fetch resource", 1)[1].split(":")[0].strip()
                for ln in (r.stderr or "").splitlines()
                if "Could not fetch resource" in ln]
    if perdidas:
        return True, "FIGURAS: " + "; ".join(perdidas[:6])
    return True, ""


#: Rutas de recurso en Markdown: `![pie](ruta)`, `[texto](ruta)` y los
#: `{{< include ruta >}}` de Quarto.
RE_RUTA_MD = re.compile(r"(!?\[[^\]]*\]\()([^)\s]+)")
RE_INCLUDE_QMD = re.compile(r"(\{\{<\s*include\s+)([^>\s]+)")


def _reubicar_rutas(texto: str, desde: pathlib.Path, hacia: pathlib.Path) -> str:
    """Reescribe las rutas relativas de un .qmd que cambia de carpeta.

    Solo toca lo que es relativo y apunta a un archivo que existe: una URL, una
    ruta absoluta o un enlace a algo inexistente se dejan como están, porque
    cambiarlos sería adivinar.
    """
    if desde.resolve() == hacia.resolve():
        return texto

    def nueva(ruta: str) -> str:
        if re.match(r"^[a-z]+:", ruta, re.I) or ruta.startswith(("/", "#")):
            return ruta
        destino = (desde / ruta).resolve()
        if not destino.exists():
            return ruta
        try:
            return os.path.relpath(destino, hacia.resolve()).replace("\\", "/")
        except ValueError:          # unidades distintas en Windows
            return ruta

    texto = RE_RUTA_MD.sub(lambda m: m.group(1) + nueva(m.group(2)), texto)
    return RE_INCLUDE_QMD.sub(lambda m: m.group(1) + nueva(m.group(2)), texto)


#: Longitud a partir de la cual una ruta de Windows deja de funcionar con las
#: llamadas normales del sistema. Vive en `rutas`; aquí queda el nombre con el
#: que lo citan los mensajes y las pruebas.
MAX_RUTA_WINDOWS = rutas.MAXIMO


def _mover(origen: pathlib.Path, destino: pathlib.Path) -> None:
    """Mueve un archivo, aguantando rutas largas en Windows."""
    try:
        origen.replace(destino)
    except OSError:
        if os.name != "nt":
            raise
        os.replace(rutas.larga(origen), rutas.larga(destino))


def _consejo_ruta_larga(destino: pathlib.Path) -> str:
    largo = len(str(destino.resolve()))
    if largo < MAX_RUTA_WINDOWS:
        return ""
    return (f"La ruta completa tiene {largo} caracteres y Windows admite "
            f"{MAX_RUTA_WINDOWS}. Mueve la carpeta de tu tesis a una ruta más "
            f"corta, por ejemplo C:\\tesis, o acorta el título.")


def comprobar_rutas(destinos: list[pathlib.Path], inf: Informe) -> bool:
    """¿Caben estas rutas en Windows? Se comprueba ANTES de generar nada.

    Los nombres que exige el art. 13 llevan apellidos, nombres, el título
    completo y el año, así que son largos por diseño. Si además la tesis está
    en una carpeta honda, se cruza el límite de 260 caracteres de Windows.

    Lo que falla entonces no es solo mover el archivo: tampoco lo abre
    python-docx para poner la carátula, ni LibreOffice para convertirlo a PDF.
    Y el error que da Windows —«el sistema no puede encontrar la ruta
    especificada»— no se parece en nada a la causa.

    Por eso se avisa antes, con el número y con qué hacer, en lugar de dejar al
    tesista con un documento a medio generar.
    """
    if os.name != "nt":
        return True
    largas = [(d, len(str(d.resolve()))) for d in destinos
              if len(str(d.resolve())) >= MAX_RUTA_WINDOWS]
    for ruta, largo in largas:
        inf.error("T28", f"la ruta de «{ruta.name}» tendría {largo} caracteres y "
                         f"Windows admite {MAX_RUTA_WINDOWS}. No se genera nada, "
                         f"para no dejarlo a medias. Mueve la carpeta de tu tesis "
                         f"a una ruta más corta —por ejemplo C:\\tesis— y vuelve "
                         f"a ejecutarlo")
    return not largas


def _bib(raiz: pathlib.Path) -> str:
    """`referencias.bib`, si existe. Sin él, las citas saldrían sin resolver."""
    return "referencias.bib" if (raiz / "referencias.bib").exists() else ""


def _csl(raiz: pathlib.Path) -> str:
    """El estilo de citas del perfil activo, leído del `_quarto.yml` generado.

    El .qmd combinado no hereda el `_quarto.yml` del proyecto, así que el CSL
    se declara en su encabezado igual que `lang`. Sin él, pandoc aplica su
    estilo por defecto (Chicago autor-año) en los documentos entregados.

    Se lee del `_quarto.yml` porque es el archivo que `activar_perfil` genera
    desde el perfil, y que ya comprobó que el .csl existe.
    """
    conf = raiz / "_quarto.yml"
    if not conf.exists():
        return ""
    try:
        datos = yaml.safe_load(conf.read_text(encoding="utf8")) or {}
    except yaml.YAMLError:
        return ""
    csl = str(datos.get("csl") or "").strip()
    return csl if csl and (raiz / csl).exists() else ""


def _limpiar_residuos(combinado: pathlib.Path) -> None:
    """Borra el .qmd temporal y lo que Quarto deja a su lado."""
    combinado.unlink(missing_ok=True)
    combinado.with_suffix(".quarto_ipynb").unlink(missing_ok=True)
    adjuntos = combinado.with_name(combinado.stem + "_files")
    if adjuntos.is_dir():
        shutil.rmtree(adjuntos, ignore_errors=True)


def _correr_quarto(cmd: list[str], raiz: pathlib.Path):
    """Corre Quarto y, si el kernel muere, limpia su caché y reintenta una vez.

    «Kernel died before replying to kernel_info» no dice nada del documento:
    es estado viejo en `.quarto/` o en `_freeze/` de un render anterior que se
    interrumpió. Pasa, y dejar al tesista con ese mensaje es abandonarlo.
    """
    r = subprocess.run(cmd, capture_output=True, text=True, cwd=str(raiz),
                       timeout=900)
    if r.returncode == 0:
        return r
    salida = (r.stderr or "") + (r.stdout or "")
    if "Kernel died" not in salida and "kernel_info" not in salida:
        return r

    for cache in (".quarto", "_freeze"):
        shutil.rmtree(raiz / cache, ignore_errors=True)
    return subprocess.run(cmd, capture_output=True, text=True, cwd=str(raiz),
                          timeout=900)


def _motivo_util(stderr: str | None, stdout: str | None) -> str:
    """El mensaje que explica el fallo, no el rastro de pila de Quarto.

    Quarto escribe su stack de JavaScript al final de la salida, así que
    quedarse con los últimos caracteres devolvía siempre lo mismo y nunca la
    causa.
    """
    texto = ((stderr or "") + "\n" + (stdout or "")).strip()
    if not texto:
        return ""
    lineas = [ln for ln in texto.splitlines()
              if ln.strip() and not ln.lstrip().startswith("at ")]
    for i, ln in enumerate(lineas):
        if "ERROR" in ln.upper():
            return " | ".join(lineas[i:i + 6])[:500]
    return " | ".join(lineas[-6:])[:500]


def _clasificar_capitulos(raiz: pathlib.Path) -> tuple[list[pathlib.Path], list[pathlib.Path]]:
    """Separa los .qmd en Desarrollo y Protocolo y Complemento (art. 13 A.2)."""
    claves_complemento = ("caratula", "carátula", "dedicatoria", "agradecimiento",
                          "indice", "índice", "lista", "referencias", "bibliografia",
                          "bibliografía", "anexo")
    desarrollo, complemento = [], []
    for ruta in sorted((raiz / "capitulos").rglob("*.qmd")):
        nombre = ruta.stem.lower()
        (complemento if any(k in nombre for k in claves_complemento) else desarrollo).append(ruta)
    return desarrollo, complemento


def preparar(raiz: pathlib.Path = RAIZ, solo_comprobar: bool = False,
             perfil_id: str | None = None) -> Informe:
    inf = Informe("entrega", "Generación de entregables", raiz=raiz)

    pid = perfil_id or perfil_activo(raiz)
    if not pid:
        inf.error("T00", "no hay perfil activo")
        return inf
    try:
        perfil = cargar_de(raiz, pid)
    except ErrorDePerfil as e:
        inf.error("T00", f"no se pudo cargar el perfil: {e}")
        return inf
    inf.anotar(perfil=pid, estado_del_perfil=perfil.estado)

    ruta_datos = raiz / "entregables" / "datos_de_entrega.yaml"
    if not ruta_datos.exists():
        ruta_datos.parent.mkdir(parents=True, exist_ok=True)
        ruta_datos.write_text(PLANTILLA_DATOS, encoding="utf8")
        inf.error("T07", f"no existía {ruta_datos.relative_to(raiz)}. Se creó una "
                         "plantilla: complétala y vuelve a ejecutar")
        return inf

    datos = cargar_datos(ruta_datos)
    autores = comprobar_datos(datos, inf)
    if inf.hay_errores:
        return inf

    titulo = str(datos["titulo"]).strip()
    anio = int(datos["anio"])
    nombres = {
        "desarrollo": nombre_archivo(autores, "Desarrollo", titulo, anio),
        "complemento": nombre_archivo(autores, "Protocolo y Complemento", titulo, anio),
    }
    inf.anotar(archivo_desarrollo=nombres["desarrollo"],
               archivo_complemento=nombres["complemento"])

    for etiqueta, nombre in nombres.items():
        if len(nombre) > 200:
            inf.aviso("T08", f"el nombre del archivo de {etiqueta} tiene {len(nombre)} "
                             "caracteres; en Windows puede dar problemas al copiarlo a "
                             "rutas largas")

    desarrollo, complemento = _clasificar_capitulos(raiz)
    inf.anotar(qmd_en_desarrollo=len(desarrollo), qmd_en_complemento=len(complemento))
    if not desarrollo:
        inf.error("T09", "no hay capítulos en `capitulos/` para el archivo de Desarrollo")

    # Documentos del expediente: se generan siempre, también en modo comprobación
    dir_documento = raiz / "entregables" / "DOCUMENTO"
    dir_documento.mkdir(parents=True, exist_ok=True)
    (dir_documento / "declaracion_jurada_pids.md").write_text(
        declaracion_jurada_pids(autores, datos), encoding="utf8")
    (dir_documento / "declaracion_uso_de_ia.md").write_text(
        declaracion_uso_de_ia(raiz, autores, datos), encoding="utf8")
    (dir_documento / "checklist_expediente.md").write_text(
        checklist_expediente(perfil, autores, {**datos, "_raiz": raiz}),
        encoding="utf8")
    (dir_documento / "autorizacion_cybertesis.md").write_text(
        autorizacion_cybertesis(autores, datos), encoding="utf8")
    (dir_documento / "metadata_complementaria.md").write_text(
        metadata_complementaria(autores, datos), encoding="utf8")
    (dir_documento / "solicitud_postergacion.md").write_text(
        solicitud_postergacion(autores, datos), encoding="utf8")
    (dir_documento / "declaracion_de_originalidad.md").write_text(
        declaracion_de_originalidad(autores, datos), encoding="utf8")
    _metadata_en_word(dir_documento, autores, datos, inf)
    inf.info("T20", f"documentos del expediente escritos en "
                    f"{dir_documento.relative_to(raiz)}")

    # El ORCID del asesor es obligatorio para el depósito (RR N.° 06893-R-18) y
    # no figura en la información complementaria del Anexo 2A, así que es el que
    # más veces falta: sin él, Biblioteca no deposita el documento.
    asesor = datos.get("asesor") or {}
    obligatorios = {
        "el ORCID del asesor": asesor.get("orcid"),
        "el DNI del asesor": asesor.get("dni"),
        "el DNI del autor": autores[0].dni if autores else None,
    }
    faltan = [nombre for nombre, valor in obligatorios.items() if not valor]
    if faltan:
        inf.aviso("T21", f"para el depósito en Cybertesis falta {', '.join(faltan)}. "
                         f"El Anexo 4 vigente —el que fijó la RR N.° 00584-R-19— los "
                         f"marca los tres como obligatorios, y ninguno aparece en la "
                         f"información complementaria del Anexo 2A. Complétalos en "
                         f"`entregables/datos_de_entrega.yaml`")

    if solo_comprobar or inf.hay_errores:
        return inf

    referencia = raiz / "plantillas" / pid / "referencia.docx"
    if not referencia.exists():
        inf.error("T14", f"falta la plantilla {referencia.relative_to(raiz)}. Genérala "
                         f"con `uv run python -m herramientas.plantilla_docx --perfil {pid}`")
        return inf

    dir_texto = raiz / "entregables" / "TEXTO"
    dir_texto.mkdir(parents=True, exist_ok=True)

    if not comprobar_rutas([dir_texto / n for n in nombres.values()], inf):
        return inf

    posicion = "inferior_derecha"
    numeracion = (perfil.get("formato") or {}).get("numeracion") or {}
    if "superior" in str(numeracion.get("posicion", "")).lower():
        posicion = "superior_derecha"

    # El reparto lo fija el art. 13 A.2, y no es cosmético: el primer archivo
    # es el que se somete al análisis de similitud, así que todo lo que no es
    # cuerpo —carátula, bibliografía, anexos— sale de él a propósito. Si se
    # queda dentro, infla el porcentaje con texto que no escribió el tesista.
    citadas = claves_citadas(desarrollo) if _bib(raiz) else []
    ya_hay_refs = any("{#refs}" in r.read_text(encoding="utf8", errors="replace")
                      for r in complemento)
    cola_refs = ""
    if citadas and not ya_hay_refs:
        cola_refs = "# Referencias {.unnumbered}\n\n::: {#refs}\n:::"

    generados: list[pathlib.Path] = []
    for etiqueta, entradas in (("desarrollo", desarrollo), ("complemento", complemento)):
        # El de Complemento se genera siempre, aunque no haya ningún .qmd
        # suyo: lleva como mínimo la carátula y la bibliografía.
        if not entradas and etiqueta == "desarrollo":
            continue
        meta = {"suppress-bibliography": "true"} if etiqueta == "desarrollo" else {}
        cola = ""
        if etiqueta == "complemento":
            if citadas:
                meta["nocite"] = "|\n  " + ", ".join(f"@{c}" for c in citadas)
            cola = cola_refs
            if not entradas and not cola:
                inf.aviso("T30", f"«{nombres[etiqueta]}» saldría vacío: no hay "
                                 f"capítulos de anexos ni citas que listar. Se "
                                 f"genera solo con la carátula")
        destino = dir_texto / nombres[etiqueta]
        ok, motivo = _render(entradas, destino, referencia, raiz, meta, cola)
        if not ok:
            inf.error("T15", f"no se pudo generar el archivo de {etiqueta}: {motivo}")
            continue
        if motivo.startswith("FIGURAS:"):
            # Pandoc sustituye la imagen que no encuentra por su texto
            # alternativo y sigue. El .docx sale, pero sin figuras.
            inf.error("T19", f"el archivo de {etiqueta} se generó SIN estas figuras, "
                             f"porque no se encontraron: "
                             f"{motivo.removeprefix('FIGURAS:').strip()}. Revisa las "
                             f"rutas en los capítulos")
        revisar_docx(destino, inf)

        # La carátula va en el archivo de «Protocolo y Complemento»: el art. 13
        # A.2 a) describe el primer archivo como el documento «sin carátula», y
        # A.2 b) pone la carátula en el segundo. Se inserta sobre el .docx YA
        # RENDERIZADO: la plantilla de referencia solo aporta estilos, y
        # copiarla traería su contenido de muestra.
        if etiqueta == "complemento":
            try:
                for aviso in caratula.anteponer(destino, datos, perfil, "tesis", raiz):
                    inf.aviso("T27", aviso, archivo=destino)
                inf.info("T22", "carátula con el escudo de la UNMSM insertada",
                         archivo=destino)
            except Exception as e:  # noqa: BLE001
                inf.aviso("T23", f"no se pudo insertar la carátula: {e}. El resto "
                                 f"del documento está bien; la carátula se puede "
                                 f"pegar a mano", archivo=destino)
            for falta in caratula.faltantes(datos, perfil):
                inf.aviso("T24", f"la carátula va incompleta: falta {falta}")

            # Después de la carátula: la numeración busca dónde empieza el
            # cuerpo para pasar de romanos a arábigos, y la carátula es parte
            # de lo que va en romanos.
            try:
                docx_numeracion.aplicar(destino, posicion=posicion)
            except docx_numeracion.SinInicioDeCuerpo:
                inf.info("T16", "el archivo de Protocolo y Complemento no lleva cuerpo, "
                                "así que va entero en numeración romana", archivo=destino)
            except Exception as e:  # noqa: BLE001
                inf.aviso("T17", f"no se pudo aplicar la numeración: {e}", archivo=destino)
        generados.append(destino)
        inf.info("T18", f"generado {destino.relative_to(raiz)}")

    _convertir_a_pdf(generados, inf, raiz)
    return inf


def _convertir_a_pdf(docxs: list[pathlib.Path], inf: Informe,
                     raiz: pathlib.Path) -> None:
    """El PDF del depósito en Cybertesis, desde el mismo .docx de la entrega.

    Nunca bloquea: el .docx es el entregable que exige el art. 13 y se genera
    igual. El PDF es lo que pide la directiva de depósito (RR N.° 06893-R-18).
    """
    if not docxs:
        return
    hechos, problemas = pdf.convertir_varios(docxs)
    for ruta in hechos:
        inf.info("T25", f"PDF generado: {ruta.relative_to(raiz)}")
    for p in problemas:
        inf.aviso("T26", f"no se generó el PDF. El .docx, que es el entregable "
                         f"del art. 13, está completo. {p}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--comprobar", action="store_true",
                    help="valida los datos y genera el expediente, sin renderizar")
    ap.add_argument("--perfil")
    ap.add_argument("--raiz", type=pathlib.Path, default=RAIZ,
                    help="raíz del proyecto; por defecto, este repositorio")
    ap.add_argument("--sin-informe", action="store_true")
    ap.add_argument("--detalle", action="store_true")
    args = ap.parse_args(argv)

    inf = preparar(raiz=args.raiz, solo_comprobar=args.comprobar,
                   perfil_id=args.perfil)
    print(inf.volcar(en_disco=not args.sin_informe, mostrar_info=args.detalle))
    return inf.codigo_salida()


if __name__ == "__main__":
    sys.exit(main())
