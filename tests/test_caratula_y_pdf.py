"""Pruebas de la carátula, el escudo y la conversión a PDF.

La que más importa: **el escudo tiene que acabar en la primera página**. Es la
clase de fallo que no da error —el .docx se abre, `inline_shapes` dice que hay
una imagen— y que solo se ve mirando el documento. Pasó dos veces al
construirlo, por dos motivos distintos, y las dos están cubiertas aquí.
"""
from __future__ import annotations

import pathlib

import docx
import pytest
from docx.oxml.ns import qn

from herramientas import caratula, marca, pdf
from herramientas.perfiles import cargar
from tests.conftest import escribir_perfil

RAIZ = pathlib.Path(__file__).resolve().parent.parent

DATOS = {
    "titulo": "Una tesis de prueba sobre algo medible",
    "anio": 2026,
    "facultad": "Facultad de Prueba",
    "escuela": "Escuela Profesional de Prueba",
    "titulo_profesional": "Licenciado en Pruebas",
    "autores": [{"nombres": "Luis Alberto", "apellidos": "Rodas Díaz"}],
    "asesor": {"nombres": "Ana María", "apellidos": "Quispe Vargas",
               "grado": "Dra.", "genero": "f"},
}


@pytest.fixture
def perfil():
    return cargar("unmsm-general")


# --- el escudo ------------------------------------------------------------

def test_el_escudo_existe_y_no_trae_la_palabra_unmsm():
    """Las carátulas publicadas usan el blasón solo; el nombre va como texto."""
    escudo = RAIZ / caratula.ESCUDO
    assert escudo.exists(), "falta el escudo; genéralo con herramientas.marca"
    from PIL import Image
    im = Image.open(escudo)
    ancho, alto = im.size
    assert alto > 500, "el recorte es demasiado bajo para ser el escudo"
    # El logo completo es bastante más alto que ancho porque lleva el texto
    # debajo; el escudo solo es casi cuadrado.
    assert alto / ancho < 1.4, ("la proporción sugiere que el recorte se llevó "
                                "también el texto del logo")


def test_el_escudo_se_puede_regenerar_desde_el_logo(tmp_path):
    """El recorte es reproducible, no un PNG que alguien hizo una vez."""
    origen = marca.origen(RAIZ)
    assert origen is not None, "no está el logo de origen en el repositorio"

    destino = tmp_path / "plantillas" / "marca"
    destino.mkdir(parents=True)
    import shutil
    shutil.copy2(origen, destino / origen.name)

    ruta_escudo, _ = marca.preparar(tmp_path)
    assert ruta_escudo.exists()

    from PIL import Image
    nuevo = Image.open(ruta_escudo).size
    actual = Image.open(RAIZ / caratula.ESCUDO).size
    assert nuevo == actual, ("regenerar el escudo da otro recorte: el del "
                             "repositorio no se hizo con esta herramienta")


# --- la carátula ----------------------------------------------------------

def test_la_caratula_sigue_el_orden_de_las_tesis_publicadas(perfil):
    texto = caratula.texto_plano(DATOS, perfil, "tesis")
    lineas = [ln for ln in texto.splitlines() if ln.strip()]
    esperado = ["[escudo de la UNMSM]",
                "Universidad Nacional Mayor de San Marcos",
                "Universidad del Perú. Decana de América",
                "Facultad de Prueba",
                "Escuela Profesional de Prueba"]
    assert lineas[:5] == esperado
    assert lineas[-2:] == ["Lima, Perú", "2026"]


def test_los_apellidos_van_en_versales():
    """Sirve para que no haya duda de cuál es el apellido y cuál el nombre."""
    assert caratula.nombre_con_apellidos_en_versales(
        "Kathia Tereza", "Basurto Salazar") == "Kathia Tereza BASURTO SALAZAR"


def test_el_rotulo_del_proyecto_lo_distingue_de_la_tesis(perfil):
    assert "PROYECTO DE TESIS" in caratula.texto_plano(DATOS, perfil, "proyecto")
    tesis = caratula.texto_plano(DATOS, perfil, "tesis")
    assert "PROYECTO" not in tesis
    assert "TESIS" in tesis


def test_el_rotulo_del_asesor_se_declina_solo_si_se_declara(perfil):
    texto = caratula.texto_plano(DATOS, perfil, "tesis")
    assert "ASESORA" in texto

    sin_genero = {**DATOS, "asesor": {"nombres_y_apellidos": "Mauro Mamani"}}
    texto = caratula.texto_plano(sin_genero, perfil, "tesis")
    assert "ASESOR" in texto and "ASESORA" not in texto


def test_un_tipo_de_caratula_desconocido_falla_claro(perfil):
    doc = docx.Document()
    with pytest.raises(ValueError, match="desconocido"):
        caratula.construir(doc, DATOS, perfil, "memoria")


def test_los_faltantes_se_nombran_en_palabras_del_tesista(perfil):
    faltan = caratula.faltantes({}, perfil)
    assert any("título de la tesis" in f for f in faltan)
    assert any("escuela" in f for f in faltan)
    assert caratula.faltantes(DATOS, perfil) == []


def test_un_proyecto_sin_escudo_usa_el_del_repositorio(tmp_path, perfil):
    """Una copia del proyecto a la que le falte `plantillas/marca/` no puede
    quedarse sin carátula: es el mismo archivo, está en las herramientas."""
    assert caratula.buscar_escudo(tmp_path) == RAIZ / caratula.ESCUDO
    doc = docx.Document()
    assert caratula.construir(doc, DATOS, perfil, "tesis", raiz=tmp_path) == []


def test_la_caratula_avisa_si_no_encuentra_el_escudo(tmp_path, perfil, monkeypatch):
    """Una carátula sin escudo que no avisa es una carátula que se devuelve."""
    monkeypatch.setattr(caratula, "buscar_escudo", lambda raiz=None: None)
    doc = docx.Document()
    avisos = caratula.construir(doc, DATOS, perfil, "tesis", raiz=tmp_path)
    assert avisos and "escudo" in avisos[0]


# --- anteponerla a un documento existente ---------------------------------

def _documento_con_cuerpo(ruta: pathlib.Path) -> None:
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph("Primer párrafo del cuerpo de la tesis.")
    doc.save(str(ruta))


def test_la_caratula_queda_delante_del_cuerpo(tmp_path, perfil):
    ruta = tmp_path / "tesis.docx"
    _documento_con_cuerpo(ruta)

    caratula.anteponer(ruta, DATOS, perfil, "tesis", raiz=RAIZ)

    doc = docx.Document(str(ruta))
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    assert textos[0] == "Universidad Nacional Mayor de San Marcos"
    assert "Introducción" in textos
    assert textos.index("Introducción") > textos.index("Lima, Perú")


def test_el_escudo_llega_al_primer_parrafo(tmp_path, perfil):
    """Si el párrafo del escudo se queda atrás, la página sale en blanco.

    Pasó por contar «los N últimos hijos» del cuerpo: `add_paragraph` inserta
    antes del `w:sectPr` final, así que esa cuenta dejaba fuera justo el primer
    párrafo nuevo, que es el del escudo.
    """
    ruta = tmp_path / "tesis.docx"
    _documento_con_cuerpo(ruta)
    caratula.anteponer(ruta, DATOS, perfil, "tesis", raiz=RAIZ)

    doc = docx.Document(str(ruta))
    assert "<w:drawing" in doc.paragraphs[0]._p.xml, (
        "el primer párrafo no lleva el escudo")


def test_la_imagen_conserva_su_relacion(tmp_path, perfil):
    """El otro modo de perder el escudo sin que nada falle.

    Si la carátula se construye en un documento aparte y se mueven sus
    elementos XML, el párrafo viaja con una referencia a una relación que vive
    en el otro documento. El .docx se abre y la imagen no está.
    """
    import re
    import zipfile

    ruta = tmp_path / "tesis.docx"
    _documento_con_cuerpo(ruta)
    caratula.anteponer(ruta, DATOS, perfil, "tesis", raiz=RAIZ)

    with zipfile.ZipFile(ruta) as z:
        assert any(n.startswith("word/media/") for n in z.namelist()), (
            "el paquete no contiene la imagen")
        rels = z.read("word/_rels/document.xml.rels").decode("utf8")

    doc = docx.Document(str(ruta))
    embebidos = re.findall(r'r:embed="([^"]+)"', doc.element.body.xml)
    assert embebidos, "no hay ninguna imagen referenciada en el cuerpo"
    for rid in embebidos:
        assert f'Id="{rid}"' in rels, f"la relación {rid} no existe en el .docx"


def test_el_sectpr_no_se_mueve_al_principio(tmp_path, perfil):
    """Mover el `w:sectPr` al frente rompe la configuración de página."""
    ruta = tmp_path / "tesis.docx"
    _documento_con_cuerpo(ruta)
    caratula.anteponer(ruta, DATOS, perfil, "tesis", raiz=RAIZ)

    cuerpo = docx.Document(str(ruta)).element.body
    hijos = list(cuerpo)
    sect = [i for i, e in enumerate(hijos) if e.tag == qn("w:sectPr")]
    if sect:
        assert sect[0] == len(hijos) - 1, "el sectPr dejó de estar al final"


# --- conversión a PDF -----------------------------------------------------

def test_buscar_soffice_no_inventa_una_ruta(monkeypatch):
    monkeypatch.delenv("SOFFICE", raising=False)
    monkeypatch.setattr(pdf.shutil, "which", lambda _: None)
    monkeypatch.setattr(pdf, "CANDIDATOS", ("/no/existe/soffice",))
    assert pdf.buscar_soffice() is None


def test_sin_libreoffice_se_explica_como_instalarlo(monkeypatch, tmp_path):
    monkeypatch.setattr(pdf, "buscar_soffice", lambda: None)
    docx_falso = tmp_path / "x.docx"
    docx_falso.write_bytes(b"PK")
    with pytest.raises(pdf.SinConversor) as e:
        pdf.convertir(docx_falso)
    assert "winget" in str(e.value) and "brew" in str(e.value)


def test_convertir_varios_no_se_detiene_en_el_primer_fallo(monkeypatch, tmp_path):
    hechos = []

    def falso(ruta, destino=None, tiempo_max=300):
        if "malo" in pathlib.Path(ruta).name:
            raise RuntimeError("no se pudo")
        salida = tmp_path / (pathlib.Path(ruta).stem + ".pdf")
        salida.write_bytes(b"%PDF")
        hechos.append(salida)
        return salida

    monkeypatch.setattr(pdf, "convertir", falso)
    ok, problemas = pdf.convertir_varios(
        [tmp_path / "bueno1.docx", tmp_path / "malo.docx", tmp_path / "bueno2.docx"])
    assert len(ok) == 2, "un fallo no puede impedir convertir los demás"
    assert len(problemas) == 1


@pytest.mark.skipif(not pdf.disponible(), reason="no hay LibreOffice en este equipo")
def test_un_docx_se_convierte_a_pdf_de_verdad(tmp_path, perfil):
    ruta = tmp_path / "doc.docx"
    _documento_con_cuerpo(ruta)
    caratula.anteponer(ruta, DATOS, perfil, "tesis", raiz=RAIZ)

    salida = pdf.convertir(ruta)
    assert salida.exists() and salida.suffix == ".pdf"
    assert salida.read_bytes()[:4] == b"%PDF"
    assert salida.stat().st_size > 5000, "el PDF salió sospechosamente pequeño"


# --- que el perfil no se quede sin plantilla ------------------------------

def test_un_proyecto_de_prueba_tiene_lo_necesario_para_la_caratula(proyecto):
    escribir_perfil(proyecto)
    from herramientas.perfiles import cargar_de
    p = cargar_de(proyecto, "prueba")
    assert caratula.faltantes(DATOS, p) == []
