"""Pruebas del proyecto de tesis como entregable y del render de la entrega.

El proyecto de tesis es el documento que la escuela registra y aprueba, y hasta
la revisión final se quedaba en Markdown. Ahora sale en .docx y PDF, con su
carátula.
"""
from __future__ import annotations

import pathlib

import pytest

from herramientas import entregar, entregar_proyecto
from herramientas.perfiles import cargar
from tests.conftest import escribir_perfil

RAIZ = pathlib.Path(__file__).resolve().parent.parent

DATOS_YAML = """\
titulo: "Una tesis de prueba"
anio: 2026
autores:
  - apellidos: "Rodas Díaz"
    nombres: "Luis Alberto"
    dni: "12345678"
    objetivos: []
asesor:
  nombres: "Ana María"
  apellidos: "Quispe Vargas"
  grado: "Dra."
  nombres_y_apellidos: "Ana María Quispe Vargas"
  correo: "a@unmsm.edu.pe"
facultad: "Facultad de Prueba"
escuela: "Escuela Profesional de Prueba"
titulo_profesional: "Licenciado en Pruebas"
"""


def _proyecto_listo(raiz: pathlib.Path) -> None:
    escribir_perfil(raiz)
    (raiz / "proyecto").mkdir(parents=True, exist_ok=True)
    (raiz / "proyecto" / "01-planteamiento.qmd").write_text(
        "# Introducción\n\nTexto.\n\n# Metodología\n\nTexto.\n", encoding="utf8")
    (raiz / "entregables").mkdir(parents=True, exist_ok=True)
    (raiz / "entregables" / "datos_de_entrega.yaml").write_text(
        DATOS_YAML, encoding="utf8")


# --- dónde vive el proyecto ----------------------------------------------

def test_el_proyecto_vive_en_su_propia_carpeta(proyecto):
    """`proyecto/` y `capitulos/` son dos documentos, no uno."""
    _proyecto_listo(proyecto)
    (proyecto / "capitulos").mkdir(exist_ok=True)
    (proyecto / "capitulos" / "01-intro.qmd").write_text("# Introducción\n",
                                                         encoding="utf8")

    hallados = entregar_proyecto.capitulos_del_proyecto(proyecto)
    assert [r.name for r in hallados] == ["01-planteamiento.qmd"]


def test_sin_proyecto_lo_dice_y_no_revienta(proyecto):
    escribir_perfil(proyecto)
    inf = entregar_proyecto.preparar(proyecto, solo_comprobar=True)
    assert inf.hay_errores
    assert any("proyecto/" in h.mensaje for h in inf.errores)


# --- qué secciones se exigen ---------------------------------------------

def test_la_caratula_no_se_reclama_como_seccion_faltante():
    """La «Primera página» la genera la herramienta; pedirla sería absurdo."""
    perfil = cargar("economia-fce")
    esperadas = entregar_proyecto.secciones_esperadas(perfil)
    assert esperadas, "economia-fce sí define secciones de proyecto"
    assert not any(t.strip().lower() in entregar_proyecto.LA_PONE_LA_CARATULA
                   for t in esperadas)


def test_una_seccion_ausente_es_aviso_y_no_bloquea(proyecto):
    """El proyecto se genera incompleto y se completa; no se detiene."""
    _proyecto_listo(proyecto)
    perfil = cargar("unmsm-general")
    from herramientas.informe import Informe
    inf = Informe("p", "p", raiz=proyecto)
    entregar_proyecto.revisar_secciones(proyecto, perfil, inf)
    assert not inf.hay_errores


# --- nada bloquea por falta del asesor ------------------------------------

def test_sin_datos_del_asesor_avisa_pero_sigue(proyecto):
    """El tesista avanza aunque todavía no tenga asesor designado."""
    _proyecto_listo(proyecto)
    datos = (proyecto / "entregables" / "datos_de_entrega.yaml").read_text(
        encoding="utf8")
    sin_asesor = datos.split("asesor:")[0] + 'facultad: "Facultad de Prueba"\n'
    (proyecto / "entregables" / "datos_de_entrega.yaml").write_text(
        sin_asesor, encoding="utf8")

    inf = entregar_proyecto.preparar(proyecto, solo_comprobar=True)
    assert not inf.hay_errores, [h.como_linea() for h in inf.errores]
    assert any("asesor" in h.mensaje for h in inf.avisos)


# --- el render de la entrega ----------------------------------------------

def test_las_rutas_relativas_se_reubican_al_combinar(tmp_path):
    """Un capítulo que cambia de carpeta se lleva sus figuras consigo."""
    raiz = tmp_path
    (raiz / "capitulos").mkdir()
    (raiz / "resultados" / "figuras").mkdir(parents=True)
    (raiz / "resultados" / "figuras" / "g.png").write_bytes(b"\x89PNG")

    texto = "![Pie](../resultados/figuras/g.png)\n\n[enlace](https://x.pe)\n"
    nuevo = entregar._reubicar_rutas(texto, raiz / "capitulos", raiz)

    assert "(resultados/figuras/g.png)" in nuevo
    assert "https://x.pe" in nuevo, "una URL no es una ruta relativa"


def test_no_se_tocan_las_rutas_que_no_existen(tmp_path):
    """Reescribir una ruta rota sería adivinar."""
    (tmp_path / "capitulos").mkdir()
    texto = "![Pie](../resultados/no_esta.png)\n"
    assert entregar._reubicar_rutas(texto, tmp_path / "capitulos", tmp_path) == texto


def test_el_motivo_del_fallo_no_es_el_rastro_de_pila():
    """Quarto termina su salida con su stack de JavaScript.

    Quedarse con los últimos caracteres devolvía siempre eso y nunca la causa.
    """
    stderr = ("ERROR: ConnectionReset en el kernel\n"
              "    at Command.actionHandler (quarto.js:142828:11)\n"
              "    at async quarto4 (quarto.js:188195:5)\n")
    motivo = entregar._motivo_util(stderr, "")
    assert "ConnectionReset" in motivo
    assert "quarto.js" not in motivo


def test_una_figura_envuelta_en_tabla_no_cuenta_como_tabla_rasterizada(tmp_path):
    """Pandoc mete cada figura con su leyenda en una tabla de 1x1.

    Contarlas como infracción del art. 13 A.6 bloqueaba la entrega de cualquier
    tesis que tuviera un gráfico.
    """
    import docx

    from herramientas.informe import Informe

    ruta = tmp_path / "d.docx"
    doc = docx.Document()
    doc.add_paragraph("Texto")
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run().add_picture(
        str(RAIZ / "plantillas" / "marca" / "escudo_unmsm.png"), width=docx.shared.Cm(2))
    doc.save(str(ruta))

    inf = Informe("t", "t", raiz=tmp_path)
    entregar.revisar_docx(ruta, inf)
    assert not any(h.codigo == "T11" for h in inf.errores)


def test_una_tabla_de_datos_rasterizada_sigue_siendo_un_error(tmp_path):
    """El filtro anterior no puede tapar lo que el art. 13 A.6 sí prohíbe."""
    import docx

    from herramientas.informe import Informe

    ruta = tmp_path / "d.docx"
    doc = docx.Document()
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].paragraphs[0].add_run().add_picture(
        str(RAIZ / "plantillas" / "marca" / "escudo_unmsm.png"), width=docx.shared.Cm(2))
    doc.save(str(ruta))

    inf = Informe("t", "t", raiz=tmp_path)
    entregar.revisar_docx(ruta, inf)
    assert any(h.codigo == "T11" for h in inf.errores)


# --- el flujo completo, si hay con qué ------------------------------------

@pytest.mark.skipif(entregar._quarto() is None, reason="no hay Quarto en este equipo")
def test_el_proyecto_se_genera_entero(proyecto):
    from herramientas import plantilla_docx
    from herramientas.perfiles import cargar_de
    _proyecto_listo(proyecto)
    plantilla_docx.construir(cargar_de(proyecto, "prueba"), raiz=proyecto)

    inf = entregar_proyecto.preparar(proyecto)
    assert not inf.hay_errores, [h.como_linea() for h in inf.errores]

    generados = list((proyecto / "entregables" / "PROYECTO").glob("*.docx"))
    assert generados, "no se generó el .docx del proyecto"

    import docx
    doc = docx.Document(str(generados[0]))
    assert "<w:drawing" in doc.paragraphs[0]._p.xml, "falta el escudo"
    textos = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "PROYECTO DE TESIS" in textos


# --- el límite de rutas de Windows ----------------------------------------

def test_una_ruta_demasiado_larga_se_avisa_antes_de_generar(tmp_path, monkeypatch):
    """Los nombres del art. 13 son largos por diseño.

    Si la tesis está en una carpeta honda se cruzan los 260 caracteres de
    Windows, y entonces no falla solo mover el archivo: tampoco lo abre
    python-docx para poner la carátula ni LibreOffice para convertirlo. El
    error de Windows no se parece en nada a la causa, así que se avisa antes.
    """
    from herramientas.informe import Informe

    monkeypatch.setattr(entregar.os, "name", "nt")
    inf = Informe("t", "t", raiz=tmp_path)
    larga = tmp_path / ("x" * 150) / ("y" * 150) / "tesis.docx"

    assert entregar.comprobar_rutas([larga], inf) is False
    assert any(h.codigo == "T28" for h in inf.errores)
    mensaje = inf.errores[0].mensaje
    assert "260" in mensaje
    assert "ruta más corta" in mensaje, "tiene que decir qué hacer"


def test_una_ruta_normal_no_molesta(tmp_path, monkeypatch):
    from herramientas.informe import Informe
    monkeypatch.setattr(entregar.os, "name", "nt")
    inf = Informe("t", "t", raiz=tmp_path)
    assert entregar.comprobar_rutas([tmp_path / "tesis.docx"], inf) is True
    assert not inf.hay_errores


def test_fuera_de_windows_no_hay_limite(tmp_path, monkeypatch):
    from herramientas.informe import Informe
    monkeypatch.setattr(entregar.os, "name", "posix")
    inf = Informe("t", "t", raiz=tmp_path)
    larga = tmp_path / ("x" * 150) / ("y" * 150) / "tesis.docx"
    assert entregar.comprobar_rutas([larga], inf) is True
    assert not inf.hay_errores


# --- el esquema por enfoque del proyecto ----------------------------------

def test_medicina_trae_un_esquema_de_proyecto_por_enfoque():
    """El anexo C es para el cuantitativo y el D para el cualitativo.

    Medicina los declara en `proyecto.modelos`, no en `proyecto.secciones`, y
    hasta la revisión de publicación nadie los leía: el proyecto cualitativo
    se generaba sin comprobar ni una sección del anexo que lo rige.
    """
    perfil = cargar("medicina-fm")
    assert entregar_proyecto.secciones_esperadas(perfil) == [], \
        "sin enfoque no hay esquema que comprobar"

    cual = entregar_proyecto.secciones_esperadas(perfil, "cualitativo")
    assert "CAPITULO I" in cual
    assert any(t.startswith("Referencias") for t in cual)
    assert "Anexos" in cual


def test_el_indice_no_se_reclama_como_seccion_del_qmd():
    """El índice de un .docx es un campo de Word, no texto del .qmd."""
    perfil = cargar("medicina-fm")
    esperadas = entregar_proyecto.secciones_esperadas(perfil, "cualitativo")
    assert not any(t.strip().lower() in entregar_proyecto.LO_PONE_WORD
                   for t in esperadas)
    assert any(t.strip().lower() in entregar_proyecto.LO_PONE_WORD
               for t in entregar_proyecto.titulos_del_esquema(perfil, "cualitativo")), \
        "el anexo D sí pide índice; lo que no pide es escribirlo a mano"


def test_el_enfoque_declarado_elige_el_esquema(proyecto):
    """`.enfoque` es lo que decide, no el orden de los modelos del perfil."""
    from herramientas.informe import Informe

    _proyecto_listo(proyecto)
    (proyecto / ".enfoque").write_text("cualitativo", encoding="utf8")
    perfil = cargar("medicina-fm")

    inf = Informe("p", "p", raiz=proyecto)
    entregar_proyecto.revisar_secciones(proyecto, perfil, inf, "cualitativo")
    faltan = [h.mensaje for h in inf.avisos if h.codigo == "Y06"]
    assert len(faltan) == 5, "un proyecto de dos secciones no cumple el anexo D"
    assert not inf.hay_errores, "faltar secciones avisa, no bloquea"
