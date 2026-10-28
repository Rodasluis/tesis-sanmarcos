"""Pruebas de las plantillas .docx y del _quarto.yml generado."""
from __future__ import annotations

import pathlib

import pytest
from docx import Document

from herramientas import activar_perfil as A
from herramientas import perfiles as P

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PLANTILLAS = RAIZ / "plantillas"


def ids():
    return P.listar()


@pytest.mark.parametrize("pid", ids())
def test_cada_perfil_tiene_plantilla_y_ficha_de_formato(pid):
    assert (PLANTILLAS / pid / "referencia.docx").exists(), (
        f"falta la plantilla de {pid}: "
        f"uv run python -m herramientas.plantilla_docx --perfil {pid}"
    )
    assert (PLANTILLAS / pid / "FORMATO.md").exists()


def test_letras_aplica_sus_margenes_y_su_tipografia():
    doc = Document(str(PLANTILLAS / "letras-flch" / "referencia.docx"))
    s = doc.sections[0]
    assert round(s.top_margin.cm, 1) == 3.0
    assert round(s.left_margin.cm, 1) == 3.0
    assert round(s.right_margin.cm, 1) == 2.5
    assert round(s.bottom_margin.cm, 1) == 2.5
    normal = next(st for st in doc.styles if st.name == "Normal")
    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 12


def test_derecho_aplica_interlineado_simple():
    """Su art. 24 j) exige espacio simple, contra el 1,5 habitual."""
    doc = Document(str(PLANTILLAS / "derecho-fdcp" / "referencia.docx"))
    cuerpo = next(st for st in doc.styles if st.name == "Body Text")
    espaciado = cuerpo.paragraph_format.line_spacing
    assert espaciado in (None, 1.0) or float(espaciado) == 1.0


@pytest.mark.parametrize("pid", ids())
def test_todas_las_plantillas_son_a4(pid):
    doc = Document(str(PLANTILLAS / pid / "referencia.docx"))
    s = doc.sections[0]
    assert round(s.page_width.cm, 1) == 21.0
    assert round(s.page_height.cm, 1) == 29.7


@pytest.mark.parametrize("pid", ids())
def test_la_ficha_de_formato_declara_lo_que_es_operativo(pid):
    texto = (PLANTILLAS / pid / "FORMATO.md").read_text(encoding="utf8")
    assert "## Valores con los que se generó" in texto
    assert ("valores que NO vienen de una norma" in texto
            or "No se aplicó" in texto), (
        f"{pid}: la ficha debe decir si algún valor es operativo o si todos vienen de la norma"
    )


def test_quarto_yml_apunta_a_la_plantilla_y_al_csl_del_perfil():
    perfil = P.cargar("medicina-fm")
    texto, avisos = A.construir_quarto_yml(perfil)
    assert "plantillas/medicina-fm/referencia.docx" in texto
    assert "plantillas/csl/vancouver.csl" in texto
    assert "lang: es" in texto
    assert "bibliography: referencias.bib" in texto
    assert not any("estilo de citas" in a for a in avisos)


def test_quarto_yml_avisa_cuando_el_perfil_no_fija_estilo():
    perfil = P.cargar("sociales-fccss")
    texto, avisos = A.construir_quarto_yml(perfil)
    assert "csl:" not in texto, "sin estilo decidido, no debe inventarse un csl"
    assert any("estilo de citas" in a for a in avisos)


def test_quarto_yml_avisa_si_el_perfil_no_esta_verificado():
    """El aviso aparece para cualquier perfil no verificado, no para uno concreto."""
    perfil = P.cargar("economia-fce")
    assert not perfil.verificado, "si se verificara, elige otro para esta prueba"
    _, avisos = A.construir_quarto_yml(perfil)
    assert any(perfil.estado in a for a in avisos), avisos


def test_un_csl_que_no_existe_se_avisa_y_no_se_escribe(tmp_path):
    """Un `csl:` roto no falla al activar el perfil: falla al renderizar.

    Dos perfiles apuntaban a `apa-es.csl`, que nunca existió. El aviso tiene que
    salir aquí, no media hora después en un mensaje de pandoc.
    """
    perfil = P.cargar("letras-flch")
    perfil.datos["csl"] = "plantillas/csl/no_existe.csl"
    texto, avisos = A.construir_quarto_yml(perfil, raiz=tmp_path)
    assert any("no existe" in a for a in avisos), avisos
    assert "no_existe.csl" not in texto


def test_letras_usa_apa_y_no_el_inexistente_apa_es():
    """No existe un apa-es.csl; el español se logra con apa.csl + lang: es."""
    perfil = P.cargar("letras-flch")
    texto, _ = A.construir_quarto_yml(perfil)
    assert "plantillas/csl/apa.csl" in texto
    assert "apa-es.csl" not in texto
    assert "lang: es" in texto


def test_los_csl_declarados_existen_en_disco():
    for nombre in ("apa.csl", "vancouver.csl", "locales-es-ES.xml"):
        assert (RAIZ / "plantillas" / "csl" / nombre).exists()


def test_formato_md_no_se_desfasa_del_perfil():
    """`FORMATO.md` lo genera `plantilla_docx`, pero nadie lo regeneraba.

    Tres quedaron mintiendo: `administrativas-fca` decía `hereda_general`
    cuando ya era `verificado`, y `economia-fce` y `educacion-fe` seguían con
    los estados intermedios que el tesista eliminó. El tesista lee ese archivo
    para saber de dónde sale el formato de su plantilla.
    """
    import re

    import yaml

    desfasados = []
    for f in sorted((RAIZ / "plantillas").glob("*/FORMATO.md")):
        m = re.search(r"Estado del perfil: `([^`]+)`", f.read_text(encoding="utf8"))
        perfil = RAIZ / "perfiles" / f"{f.parent.name}.yaml"
        if not (m and perfil.exists()):
            continue
        real = (yaml.safe_load(perfil.read_text(encoding="utf8")) or {}).get("estado")
        if m.group(1) != real:
            desfasados.append(f"{f.parent.name}: dice {m.group(1)!r}, es {real!r}")
    assert not desfasados, (
        "regenera con `uv run python -m herramientas.plantilla_docx --todos`: "
        + "; ".join(desfasados))
