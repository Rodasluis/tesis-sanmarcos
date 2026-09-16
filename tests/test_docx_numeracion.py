"""Pruebas de la numeración de páginas: romanos en preliminares, arábigos en el cuerpo.

Es el requisito de formato más caro de implementar (Anexo 2A de la RR 00744-R-20)
y el que más fácilmente se rompe en silencio, así que se prueba sobre documentos
reales construidos con python-docx.
"""
from __future__ import annotations

import pathlib

import pytest
from docx import Document
from docx.oxml.ns import qn

from herramientas import docx_numeracion as N


def _documento(tmp_path: pathlib.Path, inicio: str = "INTRODUCCIÓN") -> pathlib.Path:
    doc = Document()
    doc.add_heading("Tesis de prueba", 0)
    for titulo, cuerpo in (("Dedicatoria", "A quien corresponda."),
                           ("Resumen", "Resumen en español."),
                           ("Abstract", "Abstract in English.")):
        doc.add_heading(titulo, 1)
        doc.add_paragraph(cuerpo)
    doc.add_heading(inicio, 1)
    doc.add_paragraph("Aquí empieza el cuerpo.")
    doc.add_heading("REVISIÓN DE LITERATURA", 1)
    doc.add_paragraph("Segundo capítulo.")
    ruta = tmp_path / "tesis.docx"
    doc.save(str(ruta))
    return ruta


def test_crea_dos_secciones_con_los_formatos_correctos(tmp_path):
    ruta = _documento(tmp_path)
    N.aplicar(ruta)
    info = N.describir(ruta)
    assert info["secciones"] == 2
    prelim, cuerpo = info["detalle"]
    assert prelim["formato"] == "lowerRoman"
    assert prelim["inicio"] == "1"
    assert cuerpo["formato"] == "decimal"
    assert cuerpo["inicio"] == "1", "el cuerpo debe reiniciar en 1, no continuar"


def test_el_corte_cae_justo_antes_del_cuerpo(tmp_path):
    ruta = _documento(tmp_path)
    N.aplicar(ruta)
    doc = Document(str(ruta))
    corte = None
    for i, p in enumerate(doc.paragraphs):
        ppr = p._p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            corte = i
            break
    assert corte is not None, "no se insertó el salto de sección"
    assert doc.paragraphs[corte + 1].text.strip() == "INTRODUCCIÓN"
    textos_prelim = [p.text for p in doc.paragraphs[:corte + 1]]
    assert "Abstract in English." in textos_prelim
    assert "Aquí empieza el cuerpo." not in textos_prelim


def test_pone_el_numero_en_el_pie_por_defecto(tmp_path):
    ruta = _documento(tmp_path)
    N.aplicar(ruta)
    for d in N.describir(ruta)["detalle"]:
        assert d["campo_page_en_pie"], "falta el campo PAGE en el pie"


def test_puede_numerar_en_el_encabezado(tmp_path):
    """Letras exige el número en la esquina superior derecha."""
    ruta = _documento(tmp_path)
    N.aplicar(ruta, posicion="superior_derecha")
    for d in N.describir(ruta)["detalle"]:
        assert d["campo_page_en_encabezado"], "falta el campo PAGE en el encabezado"


def test_acepta_un_inicio_de_cuerpo_distinto(tmp_path):
    ruta = _documento(tmp_path, inicio="CAPÍTULO I. PROBLEMA DE INVESTIGACIÓN")
    N.aplicar(ruta, inicio_cuerpo="CAPÍTULO I")
    info = N.describir(ruta)
    assert info["secciones"] == 2


def test_falla_claro_si_no_encuentra_el_cuerpo(tmp_path):
    doc = Document()
    doc.add_heading("Solo preliminares", 1)
    doc.add_paragraph("Sin cuerpo.")
    ruta = tmp_path / "incompleto.docx"
    doc.save(str(ruta))
    with pytest.raises(N.SinInicioDeCuerpo):
        N.aplicar(ruta)


def test_falla_si_el_cuerpo_empieza_en_el_primer_parrafo(tmp_path):
    doc = Document()
    doc.add_heading("INTRODUCCIÓN", 1)
    doc.add_paragraph("Sin preliminares.")
    ruta = tmp_path / "sin_preliminares.docx"
    doc.save(str(ruta))
    with pytest.raises(N.SinInicioDeCuerpo, match="preliminares"):
        N.aplicar(ruta)


def test_no_modifica_el_original_si_se_pide_salida_aparte(tmp_path):
    ruta = _documento(tmp_path)
    antes = ruta.read_bytes()
    salida = tmp_path / "numerada.docx"
    N.aplicar(ruta, destino=salida)
    assert ruta.read_bytes() == antes
    assert N.describir(salida)["secciones"] == 2


def test_ignora_la_mencion_del_indice_y_usa_el_encabezado_real(tmp_path):
    """Si 'INTRODUCCIÓN' aparece antes como texto suelto, no debe cortar ahí."""
    doc = Document()
    doc.add_heading("Índice", 1)
    doc.add_paragraph("INTRODUCCIÓN")          # línea del índice, sin estilo de título
    doc.add_paragraph("REVISIÓN DE LITERATURA")
    doc.add_heading("Resumen", 1)
    doc.add_paragraph("Resumen.")
    doc.add_heading("INTRODUCCIÓN", 1)         # el cuerpo de verdad
    doc.add_paragraph("Cuerpo.")
    ruta = tmp_path / "con_indice.docx"
    doc.save(str(ruta))

    N.aplicar(ruta)
    d = Document(str(ruta))
    corte = next(i for i, p in enumerate(d.paragraphs)
                 if (p._p.find(qn("w:pPr")) is not None
                     and p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None))
    assert d.paragraphs[corte + 1].style.name.startswith("Heading")
    assert d.paragraphs[corte].text == "Resumen."
