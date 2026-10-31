"""El reparto en dos archivos que exige el art. 13 A.2.

No es cosmético: **el archivo de Desarrollo es el que se somete al análisis de
similitud**. La carátula y la bibliografía son texto que el tesista no redactó
y que coincide con el de cualquier otra tesis, así que la norma los saca de ahí
y los pone en el de «Protocolo y Complemento».
"""
from __future__ import annotations

import pathlib

import pytest

from herramientas import entregar
from tests.conftest import escribir_bib, escribir_capitulo, escribir_perfil

RAIZ = pathlib.Path(__file__).resolve().parent.parent


# --- qué claves se citan de verdad ----------------------------------------

def test_las_referencias_cruzadas_de_quarto_no_son_citas(tmp_path):
    """`@fig-dispersion` es «ver figura 1», no una obra.

    Si se colara en el `nocite` del archivo de Complemento, citeproc buscaría
    esa clave en el .bib, no la encontraría y la imprimiría en crudo.
    """
    f = tmp_path / "cap.qmd"
    f.write_text(
        "Como se ve en @fig-dispersion y en la @tbl-resumen, y segun "
        "@Czernich_2011 y @Maurseth_2018 [p. 14], el efecto de @eq-uno "
        "se discute en la @sec-metodos.", encoding="utf8")
    assert entregar.claves_citadas([f]) == ["Czernich_2011", "Maurseth_2018"]


def test_las_claves_salen_sin_repetir_y_en_orden(tmp_path):
    f = tmp_path / "cap.qmd"
    f.write_text("@b dice una cosa; @a otra; @b insiste.", encoding="utf8")
    assert entregar.claves_citadas([f]) == ["b", "a"]


def test_un_correo_no_es_una_cita(tmp_path):
    f = tmp_path / "cap.qmd"
    f.write_text("Escribe a tesis@unmsm.edu.pe para pedirlo.", encoding="utf8")
    assert entregar.claves_citadas([f]) == []


def test_sin_citas_no_hay_claves(tmp_path):
    f = tmp_path / "cap.qmd"
    f.write_text("Un capitulo sin una sola cita.", encoding="utf8")
    assert entregar.claves_citadas([f]) == []


# --- el flujo entero, si hay con qué renderizar ---------------------------

def _tesis_lista(raiz: pathlib.Path) -> None:
    escribir_perfil(raiz, "prueba")
    (raiz / ".perfil-activo").write_text("prueba", encoding="utf8")
    escribir_bib(raiz, """
@article{Czernich_2011,
  author = {Czernich, Nina},
  title = {Broadband Infrastructure and Economic Growth},
  journal = {The Economic Journal},
  year = {2011}
}
""".strip())
    escribir_capitulo(raiz, "01-introduccion.qmd",
                      "# Introduccion\n\nSegun @Czernich_2011, algo pasa.\n")
    (raiz / "entregables").mkdir(exist_ok=True)
    (raiz / "entregables" / "datos_de_entrega.yaml").write_text(
        "titulo: Una tesis de prueba\n"
        "titulo_profesional: Licenciado en Pruebas\n"
        "facultad: Facultad de Prueba\n"
        "escuela: Escuela de Prueba\n"
        "anio: 2026\n"
        "autores:\n"
        "  - apellidos: Rojas Diaz\n"
        "    nombres: Luis\n"
        "    dni: '12345678'\n"
        "asesor:\n"
        "  nombres_y_apellidos: Ana Quispe\n",
        encoding="utf8")


@pytest.mark.skipif(entregar._quarto() is None, reason="no hay Quarto en este equipo")
def test_la_caratula_y_la_bibliografia_van_en_el_segundo_archivo(proyecto):
    import docx

    from herramientas import plantilla_docx
    from herramientas.perfiles import cargar_de

    _tesis_lista(proyecto)
    plantilla_docx.construir(cargar_de(proyecto, "prueba"), raiz=proyecto)

    inf = entregar.preparar(proyecto)
    assert not inf.hay_errores, [h.como_linea() for h in inf.errores]

    texto = proyecto / "entregables" / "TEXTO"
    desarrollo = list(texto.glob("*Desarrollo*.docx"))
    complemento = list(texto.glob("*Protocolo y Complemento*.docx"))
    assert desarrollo, "falta el archivo de Desarrollo"
    assert complemento, (
        "falta el de Protocolo y Complemento: el art. 13 A.2 pide DOS archivos")

    def parrafos(f):
        return [p.text.strip() for p in docx.Document(str(f)).paragraphs if p.text.strip()]

    d, c = parrafos(desarrollo[0]), parrafos(complemento[0])

    # A.2 a): el primero va «sin caratula […] ni referencias bibliograficas».
    assert not any("Decana de Am" in p for p in d), "la caratula no va en Desarrollo"
    assert not any(p.startswith("Czernich, Nina") for p in d), (
        "la lista de referencias no va en Desarrollo")

    # A.2 b): el segundo lleva la caratula y la bibliografia.
    assert any("Decana de Am" in p for p in c), "falta la caratula en el Complemento"
    assert any(p.startswith("Czernich, Nina") for p in c), (
        "falta la bibliografia en el Complemento")


@pytest.mark.skipif(entregar._quarto() is None, reason="no hay Quarto en este equipo")
def test_la_cita_en_el_texto_se_queda_en_el_cuerpo(proyecto):
    """Sale la LISTA de referencias, no la cita. «(Czernich 2011)» es cuerpo."""
    import docx

    from herramientas import plantilla_docx
    from herramientas.perfiles import cargar_de

    _tesis_lista(proyecto)
    plantilla_docx.construir(cargar_de(proyecto, "prueba"), raiz=proyecto)
    entregar.preparar(proyecto)

    f = list((proyecto / "entregables" / "TEXTO").glob("*Desarrollo*.docx"))[0]
    cuerpo = " ".join(p.text for p in docx.Document(str(f)).paragraphs)
    assert "Czernich" in cuerpo, "la cita en el texto tiene que seguir ahi"


# --- el art. 13 A.6 y los callouts de Quarto ------------------------------

def _docx_con_tabla(ruta: pathlib.Path, filas: int, cols: int, emu: int | None) -> None:
    """Un .docx con una tabla que lleva —o no— una imagen de `emu` de lado."""
    import docx
    from docx.shared import Emu

    d = docx.Document()
    t = d.add_table(rows=filas, cols=cols)
    t.cell(0, 0).text = "contenido"
    if emu is not None:
        png = ruta.parent / "punto.png"
        png.write_bytes(bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
            "890000000a49444154789c6300010000050001" "0d0a2db4" "0000000049454e44ae426082"))
        t.cell(0, 0).paragraphs[0].add_run().add_picture(str(png), width=Emu(emu),
                                                         height=Emu(emu))
    d.save(str(ruta))


def _revisar(ruta: pathlib.Path):
    from herramientas.informe import Informe
    inf = Informe("prueba", "Prueba")
    entregar.revisar_docx(ruta, inf)
    return inf


def test_un_callout_de_quarto_no_es_una_tabla_fotografiada(tmp_path):
    """`::: {.callout-note}` sale como tabla de dos filas con el icono dentro.

    El icono mide 152400 EMU, que son 12 puntos. Acusar de infringir el
    art. 13 A.6 a quien uso un callout lo manda a buscar un problema que no
    existe, y era un ERROR que detenia la entrega: le pasaba al ejemplo
    cualitativo del propio repositorio.
    """
    f = tmp_path / "callout.docx"
    _docx_con_tabla(f, filas=2, cols=2, emu=152400)
    assert not [h.codigo for h in _revisar(f).errores if h.codigo == "T11"]


def test_una_tabla_de_verdad_fotografiada_si_se_detecta(tmp_path):
    """Lo que el art. 13 A.6 prohibe de verdad: la tabla pegada como imagen."""
    f = tmp_path / "rasterizada.docx"
    _docx_con_tabla(f, filas=2, cols=2, emu=4572000)  # 5 pulgadas
    assert [h.codigo for h in _revisar(f).errores if h.codigo == "T11"]


def test_la_figura_envuelta_por_pandoc_sigue_sin_contar(tmp_path):
    """Pandoc mete cada figura con su leyenda en una tabla de una celda."""
    f = tmp_path / "figura.docx"
    _docx_con_tabla(f, filas=1, cols=1, emu=4572000)
    assert not [h.codigo for h in _revisar(f).errores if h.codigo == "T11"]


def test_revisar_docx_no_revienta_con_una_tabla_sin_tblgrid(tmp_path):
    """`len(tabla.columns)` lanza InvalidXmlError si falta `<w:tblGrid>`, que
    es justo lo que traen las tablas de los callouts."""
    import docx
    from docx.oxml.ns import qn

    f = tmp_path / "sin_grid.docx"
    _docx_con_tabla(f, filas=1, cols=1, emu=152400)
    d = docx.Document(str(f))
    tbl = d.tables[0]._tbl
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    d.save(str(f))

    _revisar(f)  # no debe lanzar


def test_el_icono_de_un_callout_no_dispara_el_aviso_de_formula_fotografiada(tmp_path):
    """El icono de 12 puntos de un callout no es una formula rasterizada."""
    f = tmp_path / "callout.docx"
    _docx_con_tabla(f, filas=2, cols=2, emu=152400)
    assert not [h.codigo for h in _revisar(f).avisos if h.codigo == "T12"]


def test_una_imagen_grande_sin_ecuaciones_si_se_avisa(tmp_path):
    """El aviso tiene que seguir saltando donde puede haber una formula."""
    f = tmp_path / "grande.docx"
    _docx_con_tabla(f, filas=1, cols=1, emu=4572000)
    assert [h.codigo for h in _revisar(f).avisos if h.codigo == "T12"]


# --- el estilo de citas del perfil llega al documento entregado -----------

def test_la_cabecera_del_combinado_declara_el_csl_del_perfil(tmp_path):
    """El .qmd combinado no hereda el `_quarto.yml`: el CSL va en su encabezado."""
    (tmp_path / "referencias.bib").write_text("", encoding="utf8")
    (tmp_path / "plantillas" / "csl").mkdir(parents=True)
    (tmp_path / "plantillas" / "csl" / "vancouver.csl").write_text("<style/>",
                                                                  encoding="utf8")
    (tmp_path / "_quarto.yml").write_text(
        "lang: es\ncsl: plantillas/csl/vancouver.csl\n", encoding="utf8")

    cabecera = entregar._cabecera(tmp_path)
    assert 'csl: "plantillas/csl/vancouver.csl"' in cabecera
    assert 'bibliography: "referencias.bib"' in cabecera


def test_un_csl_declarado_que_no_existe_no_se_declara(tmp_path):
    """Pandoc falla entero si el .csl no esta; mejor el estilo por defecto."""
    (tmp_path / "_quarto.yml").write_text("csl: plantillas/csl/fantasma.csl\n",
                                          encoding="utf8")
    assert "csl:" not in entregar._cabecera(tmp_path)


def test_sin_quarto_yml_la_cabecera_sigue_saliendo(tmp_path):
    """Las pruebas y los proyectos a medio montar no tienen `_quarto.yml`."""
    cabecera = entregar._cabecera(tmp_path)
    assert cabecera.startswith("---")
    assert "lang: es" in cabecera
